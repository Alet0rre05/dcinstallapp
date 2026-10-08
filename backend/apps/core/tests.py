import io
import zipfile
from datetime import date, datetime, timedelta
from datetime import timezone as dt_timezone
from zoneinfo import ZoneInfo
from unittest import mock

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import IntegrityError, connection
from django.test import TestCase, override_settings
from django.test.utils import CaptureQueriesContext
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient

from .models import (
    Auditoria,
    Cliente,
    Control,
    LimiteMensajesExcedido,
    Matafuego,
    MensajeChat,
    Rol,
    TicketSoporte,
)

User = get_user_model()
HOY = timezone.localdate()


def mata(cliente, serie="A1", carga=None, ph=None):
    return Matafuego.objects.create(
        cliente=cliente,
        numero_serie=serie,
        vencimiento_carga=carga or HOY + timedelta(days=200),
        vencimiento_ph=ph or HOY + timedelta(days=200),
    )


class EstadoColorTests(TestCase):
    def setUp(self):
        self.cliente = Cliente.objects.create(nombre="ACME")
        self.user = User.objects.create_user("op", password="x")

    def control(self, m, **kw):
        return Control.objects.create(matafuego=m, usuario=self.user, **kw)

    def color(self, m):
        return Matafuego.objects.prefetch_related("controles").get(pk=m.pk).estado_color

    def test_gris_sin_controles(self):
        self.assertEqual(self.color(mata(self.cliente)), "GRIS")

    def test_gris_sin_vencimientos(self):
        m = Matafuego.objects.create(cliente=self.cliente, numero_serie="X")
        self.control(m)
        self.assertEqual(self.color(m), "GRIS")

    def test_verde(self):
        m = mata(self.cliente)
        self.control(m)
        self.assertEqual(self.color(m), "VERDE")

    def test_amarillo(self):
        m = mata(self.cliente, carga=HOY + timedelta(days=30))
        self.control(m)
        self.assertEqual(self.color(m), "AMARILLO")

    def test_rojo_un_problema(self):
        m = mata(self.cliente, carga=HOY - timedelta(days=1))
        self.control(m)
        self.assertEqual(self.color(m), "ROJO")

    def test_bordo_dos_problemas(self):
        m = mata(self.cliente, carga=HOY - timedelta(days=1))
        self.control(m, presion=Control.Presion.BAJA)
        self.assertEqual(self.color(m), "BORDO")

    def test_bordo_inaccesible_y_chapa(self):
        m = mata(self.cliente)
        self.control(m, accesible=False, chapa_baliza=False)
        self.assertEqual(self.color(m), "BORDO")

    def test_control_actualiza_padre_y_es_inmutable(self):
        m = mata(self.cliente)
        nueva = HOY + timedelta(days=400)
        c = self.control(m, ubicacion="Hall", vencimiento_carga=nueva)
        m.refresh_from_db()
        self.assertEqual(m.ubicacion, "Hall")
        self.assertEqual(m.vencimiento_carga, nueva)
        with self.assertRaises(PermissionError):
            c.save()
        with self.assertRaises(PermissionError):
            c.delete()

    def test_token_qr_inmutable(self):
        m = mata(self.cliente)
        original = m.token_qr
        import uuid

        m.token_qr = uuid.uuid4()
        m.save()
        m.refresh_from_db()
        self.assertEqual(m.token_qr, original)


class ChatTests(TestCase):
    def setUp(self):
        self.cliente = Cliente.objects.create(nombre="ACME")
        self.u = User.objects.create_user("cli", password="x")
        self.s = User.objects.create_user("sup", password="x")
        self.t = TicketSoporte.objects.create(titulo="Ayuda", cliente=self.cliente, creado_por=self.u)

    def enviar(self, user, soporte):
        MensajeChat.validar_limite(self.t, soporte)
        MensajeChat.objects.create(ticket=self.t, autor=user, es_soporte=soporte, texto="hola")

    def test_cliente_max_3_seguidos(self):
        for _ in range(3):
            self.enviar(self.u, False)
        with self.assertRaises(LimiteMensajesExcedido):
            self.enviar(self.u, False)
        self.enviar(self.s, True)  # soporte responde
        self.enviar(self.u, False)  # el cliente puede volver a escribir

    def test_soporte_max_10_seguidos(self):
        for _ in range(10):
            self.enviar(self.s, True)
        with self.assertRaises(LimiteMensajesExcedido):
            self.enviar(self.s, True)


@override_settings(AXES_ENABLED=False, SECURE_SSL_REDIRECT=False)
class ApiTests(TestCase):
    def setUp(self):
        self.api = APIClient()
        self.a = Cliente.objects.create(nombre="A")
        self.b = Cliente.objects.create(nombre="B")
        self.u = User.objects.create_user("u", password="Clave-Segura-123")
        self.u.perfil.rol = Rol.OFICINA
        self.u.perfil.save()
        self.u.perfil.clientes.add(self.a)

    @override_settings(DEBUG=True)
    def test_registro_crea_publico_sin_rol(self):
        r = self.api.post(
            "/api/auth/registro/",
            {"username": "nuevo", "email": "n@x.com", "first_name": "N", "password": "Clave-Segura-123"},
            format="json",
        )
        self.assertEqual(r.status_code, 201, r.content)
        self.assertIsNone(r.json()["user"]["rol"])
        self.assertTrue(Auditoria.objects.filter(accion="REGISTRO", usuario_nombre="nuevo").exists())

    @override_settings(DEBUG=False, TURNSTILE_SECRET_KEY="")
    def test_registro_sin_turnstile_falla_en_produccion(self):
        r = self.api.post(
            "/api/auth/registro/",
            {"username": "nuevo", "email": "n@x.com", "first_name": "N", "password": "Clave-Segura-123"},
            format="json",
        )
        self.assertEqual(r.status_code, 400)

    def test_multitenancy_matafuegos(self):
        ma, mb = mata(self.a, "1"), mata(self.b, "2")
        self.api.force_authenticate(self.u)
        ids = [x["id"] for x in self.api.get("/api/matafuegos/").json()["results"]]
        self.assertIn(ma.id, ids)
        self.assertNotIn(mb.id, ids)

    def test_qr_publico_datos_reducidos(self):
        m = mata(self.a)
        r = self.api.get(f"/api/public/qr/{m.token_qr}/")
        self.assertEqual(r.status_code, 200)
        self.assertNotIn("id", r.json())
        self.assertNotIn("token_qr", r.json())
        self.assertEqual(r.json()["cliente"], "A")

    def test_baja_logica(self):
        m = mata(self.a)
        self.api.force_authenticate(self.u)
        self.assertEqual(self.api.delete(f"/api/matafuegos/{m.id}/").status_code, 204)
        m.refresh_from_db()
        self.assertFalse(m.activo)
        self.assertEqual(self.api.get(f"/api/public/qr/{m.token_qr}/").status_code, 404)

    def test_cambio_de_datos_exige_password_actual(self):
        self.api.force_authenticate(self.u)
        r = self.api.patch("/api/auth/me/", {"first_name": "Z", "current_password": "mala"}, format="json")
        self.assertEqual(r.status_code, 400)
        r = self.api.patch(
            "/api/auth/me/", {"first_name": "Z", "current_password": "Clave-Segura-123"}, format="json"
        )
        self.assertEqual(r.status_code, 200)

    def test_chat_api_limite_y_paginacion(self):
        self.api.force_authenticate(self.u)
        for i in range(7):
            self.api.post("/api/tickets/", {"titulo": f"T{i}", "cliente": self.a.id}, format="json")
        r = self.api.get("/api/tickets/").json()
        self.assertEqual(len(r["results"]), 5)
        self.assertEqual(r["count"], 7)
        tid = r["results"][0]["id"]
        for _ in range(10):
            self.assertEqual(
                self.api.post(f"/api/tickets/{tid}/mensajes/", {"texto": "hola"}, format="json").status_code,
                201,
            )
        self.assertEqual(
            self.api.post(f"/api/tickets/{tid}/mensajes/", {"texto": "hola"}, format="json").status_code,
            429,
        )


@override_settings(AXES_ENABLED=False, SECURE_SSL_REDIRECT=False)
class RolesYEquipoTests(TestCase):
    def setUp(self):
        self.api = APIClient()
        self.a = Cliente.objects.create(nombre="A")
        self.b = Cliente.objects.create(nombre="B")
        self.admin = User.objects.create_superuser("admin", "admin@x.com", "Clave-Segura-123")
        self.oficina = self._user("ofi", "ofi@x.com", Rol.OFICINA, [self.a])
        self.operario = self._user("ope", "ope@x.com", Rol.OPERARIO, [self.a])
        self.publico = self._user("pub", "pub@x.com", None, [])

    def _user(self, username, email, rol, clientes):
        u = User.objects.create_user(username, email, "Clave-Segura-123")
        u.perfil.rol = rol
        u.perfil.save()
        u.perfil.clientes.set(clientes)
        return u

    # -- Público registrado -------------------------------------------------
    def test_publico_sin_acceso_a_nada_interno(self):
        self.api.force_authenticate(self.publico)
        self.assertEqual(self.api.get("/api/matafuegos/").status_code, 403)
        self.assertEqual(self.api.get("/api/tickets/").status_code, 403)
        self.assertEqual(self.api.get("/api/auditoria/").status_code, 403)
        self.assertEqual(self.api.get("/api/equipo/usuarios/").status_code, 403)
        r = self.api.patch("/api/auth/me/", {"first_name": "Z", "current_password": "Clave-Segura-123"}, format="json")
        self.assertEqual(r.status_code, 403)

    def test_publico_puede_ver_me_y_qr_publico(self):
        m = mata(self.a)
        self.api.force_authenticate(self.publico)
        r = self.api.get("/api/auth/me/")
        self.assertEqual(r.status_code, 200)
        self.assertIsNone(r.json()["rol"])
        self.assertEqual(self.api.get(f"/api/public/qr/{m.token_qr}/").status_code, 200)

    def test_login_deja_auditoria(self):
        r = self.api.post("/api/auth/login/", {"username": "ofi", "password": "Clave-Segura-123"}, format="json")
        self.assertEqual(r.status_code, 200, r.content)
        self.assertTrue(Auditoria.objects.filter(accion="LOGIN", usuario_nombre="ofi").exists())

    # -- Equipo ---------------------------------------------------------------
    def test_solo_admin_gestiona_equipo(self):
        for u in (self.oficina, self.operario, self.publico):
            self.api.force_authenticate(u)
            r = self.api.post("/api/equipo/asignar-rol/", {"email": "pub@x.com", "rol": "OPERARIO"}, format="json")
            self.assertEqual(r.status_code, 403)

    def test_asignar_rol_email_inexistente_no_crea_cuenta(self):
        self.api.force_authenticate(self.admin)
        antes = User.objects.count()
        r = self.api.post("/api/equipo/asignar-rol/", {"email": "nadie@x.com", "rol": "OPERARIO"}, format="json")
        self.assertEqual(r.status_code, 404)
        self.assertIn("registrarse", r.json()["detail"])
        self.assertEqual(User.objects.count(), antes)

    def test_asignar_y_revocar_rol_con_auditoria(self):
        self.api.force_authenticate(self.admin)
        r = self.api.post(
            "/api/equipo/asignar-rol/",
            {"email": "PUB@x.com", "rol": "OPERARIO", "clientes": [self.a.id]},
            format="json",
        )
        self.assertEqual(r.status_code, 200, r.content)
        self.publico.perfil.refresh_from_db()
        self.assertEqual(self.publico.perfil.rol, Rol.OPERARIO)
        ev = Auditoria.objects.get(accion="ASIGNACION_ROL")
        self.assertEqual(ev.usuario_nombre, "admin")
        self.assertIsNone(ev.antes["rol"])
        self.assertEqual(ev.despues["rol"], "OPERARIO")

        r = self.api.post("/api/equipo/revocar-rol/", {"user_id": self.publico.id}, format="json")
        self.assertEqual(r.status_code, 200, r.content)
        self.publico.perfil.refresh_from_db()
        self.assertIsNone(self.publico.perfil.rol)
        self.assertTrue(User.objects.filter(pk=self.publico.pk).exists())  # la cuenta NO se elimina
        ev = Auditoria.objects.get(accion="REVOCACION_ROL")
        self.assertEqual(ev.antes["rol"], "OPERARIO")
        self.assertIsNone(ev.despues["rol"])

        # pierde acceso de inmediato
        self.api.force_authenticate(self.publico)
        self.assertEqual(self.api.get("/api/matafuegos/").status_code, 403)

    def test_admin_no_puede_revocarse_a_si_mismo(self):
        self.api.force_authenticate(self.admin)
        r = self.api.post("/api/equipo/revocar-rol/", {"user_id": self.admin.id}, format="json")
        self.assertEqual(r.status_code, 400)

    # -- QR con edición según rol ---------------------------------------------
    def test_qr_privado_respeta_cliente_asignado(self):
        propia, ajena = mata(self.a, "P"), mata(self.b, "X")
        self.api.force_authenticate(self.operario)
        self.assertEqual(self.api.get(f"/api/matafuegos/qr/{propia.token_qr}/").status_code, 200)
        self.assertEqual(self.api.get(f"/api/matafuegos/qr/{ajena.token_qr}/").status_code, 404)
        self.api.force_authenticate(self.publico)
        self.assertEqual(self.api.get(f"/api/matafuegos/qr/{propia.token_qr}/").status_code, 403)
        self.api.force_authenticate(None)
        self.assertEqual(self.api.get(f"/api/matafuegos/qr/{propia.token_qr}/").status_code, 401)

    def test_operario_crea_control_pero_no_edita_matafuego(self):
        m = mata(self.a)
        self.api.force_authenticate(self.operario)
        r = self.api.post("/api/controles/", {"matafuego": m.id, "presion": "NORMAL"}, format="json")
        self.assertEqual(r.status_code, 201, r.content)
        r = self.api.patch(f"/api/matafuegos/{m.id}/", {"ubicacion": "X"}, format="json")
        self.assertEqual(r.status_code, 403)

    # -- Auditoría --------------------------------------------------------------
    def test_auditoria_de_matafuego_y_visibilidad_por_rol(self):
        self.api.force_authenticate(self.oficina)
        r = self.api.post(
            "/api/matafuegos/", {"cliente": self.a.id, "numero_serie": "Z1", "clase": "ABC"}, format="json"
        )
        self.assertEqual(r.status_code, 201, r.content)
        mid = r.json()["id"]
        self.api.patch(f"/api/matafuegos/{mid}/", {"ubicacion": "Hall"}, format="json")
        self.api.delete(f"/api/matafuegos/{mid}/")
        self.api.post(f"/api/matafuegos/{mid}/restaurar/")
        acciones = list(Auditoria.objects.filter(objeto="Matafuego").values_list("accion", flat=True))
        for esperado in ("CREACION", "MODIFICACION", "BAJA", "RESTAURACION"):
            self.assertIn(esperado, acciones)
        mod = Auditoria.objects.get(accion="MODIFICACION", objeto="Matafuego")
        self.assertEqual(mod.antes["ubicacion"], "")
        self.assertEqual(mod.despues["ubicacion"], "Hall")

        # evento de otro cliente: la oficina (cliente A) no lo ve; el admin sí
        Auditoria.objects.create(accion="CREACION", objeto="Matafuego", cliente_id=self.b.id, usuario_nombre="otro")
        ids_ofi = {e["cliente_id"] for e in self.api.get("/api/auditoria/").json()["results"]}
        self.assertEqual(ids_ofi, {self.a.id})
        self.api.force_authenticate(self.admin)
        ids_admin = {e["cliente_id"] for e in self.api.get("/api/auditoria/").json()["results"]}
        self.assertIn(self.b.id, ids_admin)

    def test_operario_solo_ve_su_auditoria(self):
        m = mata(self.a)
        self.api.force_authenticate(self.operario)
        self.api.post("/api/controles/", {"matafuego": m.id}, format="json")
        Auditoria.objects.create(accion="CREACION", objeto="Matafuego", cliente_id=self.a.id, usuario_nombre="otro")
        res = self.api.get("/api/auditoria/").json()["results"]
        self.assertTrue(res)
        self.assertTrue(all(e["usuario_nombre"] == "ope" for e in res))

    def test_auditoria_es_inmutable(self):
        ev = Auditoria.objects.create(accion="X", objeto="Y")
        with self.assertRaises(PermissionError):
            ev.save()
        with self.assertRaises(PermissionError):
            ev.delete()


@override_settings(AXES_ENABLED=False, SECURE_SSL_REDIRECT=False)
class ClientesPaginacionYResumenTests(TestCase):
    """V2.2 · Paso 1: filtro por cliente, paginación grande y resumen por cliente."""

    def setUp(self):
        self.api = APIClient()
        self.a = Cliente.objects.create(nombre="A")
        self.b = Cliente.objects.create(nombre="B")
        self.vacio = Cliente.objects.create(nombre="Vacío")
        self.admin = User.objects.create_superuser("admin", "admin@x.com", "Clave-Segura-123")
        self.oficina = self._user("ofi", Rol.OFICINA, [self.a, self.b])
        self.operario = self._user("ope", Rol.OPERARIO, [self.a, self.vacio])
        self.publico = self._user("pub", None, [])

    def _user(self, username, rol, clientes):
        u = User.objects.create_user(username, f"{username}@x.com", "Clave-Segura-123")
        u.perfil.rol = rol
        u.perfil.save()
        u.perfil.clientes.set(clientes)
        return u

    def control(self, m, **kw):
        return Control.objects.create(matafuego=m, usuario=self.admin, **kw)

    def series(self, resp):
        return {x["numero_serie"] for x in resp.json()["results"]}

    # -- Filtro ?cliente= y scoping ----------------------------------------
    def test_filtro_por_cliente(self):
        mata(self.a, "A1"), mata(self.a, "A2"), mata(self.b, "B1")
        self.api.force_authenticate(self.oficina)
        self.assertEqual(self.series(self.api.get(f"/api/matafuegos/?cliente={self.a.id}")), {"A1", "A2"})
        self.assertEqual(self.series(self.api.get(f"/api/matafuegos/?cliente={self.b.id}")), {"B1"})
        self.assertEqual(self.series(self.api.get("/api/matafuegos/")), {"A1", "A2", "B1"})

    def test_cliente_ajeno_da_lista_vacia(self):
        mata(self.a, "A1"), mata(self.b, "B1")
        self.api.force_authenticate(self.operario)  # tiene A y Vacío, no B
        r = self.api.get(f"/api/matafuegos/?cliente={self.b.id}")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.json()["count"], 0)
        self.assertEqual(self.series(self.api.get("/api/matafuegos/")), {"A1"})  # y sin filtro tampoco

    def test_filtro_cliente_invalido_es_400_no_500(self):
        self.api.force_authenticate(self.oficina)
        for valor in ("abc", "1;DROP", "9" * 30, "-1"):
            self.assertEqual(self.api.get(f"/api/matafuegos/?cliente={valor}").status_code, 400, valor)

    # -- Paginación -----------------------------------------------------------
    def test_paginacion_mas_de_20_equipos(self):
        Matafuego.objects.bulk_create(
            [Matafuego(cliente=self.a, numero_serie=f"S{i:03d}") for i in range(130)]
        )
        self.api.force_authenticate(self.operario)
        r1 = self.api.get(f"/api/matafuegos/?cliente={self.a.id}").json()
        self.assertEqual(r1["count"], 130)
        self.assertEqual(len(r1["results"]), 100)  # antes eran 20: los 21+ no se veían
        self.assertIsNotNone(r1["next"])
        r2 = self.api.get(r1["next"]).json()
        self.assertEqual(len(r2["results"]), 30)
        self.assertIsNone(r2["next"])
        vistos = {x["numero_serie"] for x in r1["results"] + r2["results"]}
        self.assertEqual(len(vistos), 130)  # sin repetidos ni faltantes

    def test_page_size_configurable_y_con_tope(self):
        Matafuego.objects.bulk_create(
            [Matafuego(cliente=self.a, numero_serie=f"S{i:03d}") for i in range(130)]
        )
        self.api.force_authenticate(self.operario)
        r = self.api.get("/api/matafuegos/?page_size=500").json()
        self.assertEqual(len(r["results"]), 130)
        self.assertIsNone(r["next"])
        self.assertEqual(len(self.api.get("/api/matafuegos/?page_size=25").json()["results"]), 25)
        from .views import MatafuegoPagination

        self.assertEqual(MatafuegoPagination.max_page_size, 500)

    # -- Resumen -----------------------------------------------------------------
    def armar_estados(self):
        mata(self.a, "VERDE"), mata(self.a, "AMARILLO", carga=HOY + timedelta(days=10))
        mata(self.a, "ROJO", carga=HOY - timedelta(days=1))
        mata(self.a, "BORDO", carga=HOY - timedelta(days=1), ph=HOY - timedelta(days=1))
        mata(self.a, "GRIS")
        for m in Matafuego.objects.filter(cliente=self.a).exclude(numero_serie="GRIS"):
            self.control(m)  # GRIS queda sin controles
        baja = mata(self.a, "BAJA")
        baja.activo = False
        baja.save()
        mata(self.b, "B1")

    def test_resumen_conteo_por_estado(self):
        self.armar_estados()
        self.api.force_authenticate(self.oficina)
        r = self.api.get("/api/clientes/resumen/")
        self.assertEqual(r.status_code, 200)
        por_id = {c["id"]: c for c in r.json()}
        a = por_id[self.a.id]
        self.assertEqual(a["nombre"], "A")
        self.assertEqual(a["total"], 5)  # la baja no cuenta
        self.assertEqual(
            a["por_estado"], {"BORDO": 1, "ROJO": 1, "AMARILLO": 1, "GRIS": 1, "VERDE": 1}
        )
        self.assertEqual(a["vencidos_criticos"], 2)
        self.assertEqual(a["por_vencer"], 1)
        self.assertEqual(por_id[self.b.id]["total"], 1)

    def test_resumen_respeta_clientes_asignados(self):
        self.armar_estados()
        self.api.force_authenticate(self.operario)  # A y Vacío
        res = self.api.get("/api/clientes/resumen/").json()
        self.assertEqual({c["nombre"] for c in res}, {"A", "Vacío"})
        vacio = next(c for c in res if c["nombre"] == "Vacío")
        self.assertEqual(vacio["total"], 0)  # un cliente sin equipos igual aparece
        self.assertEqual(vacio["vencidos_criticos"], 0)
        self.api.force_authenticate(self.admin)
        self.assertEqual(len(self.api.get("/api/clientes/resumen/").json()), 3)

    def test_resumen_exige_sesion_y_rol(self):
        self.api.force_authenticate(None)
        self.assertEqual(self.api.get("/api/clientes/resumen/").status_code, 401)
        self.api.force_authenticate(self.publico)
        self.assertEqual(self.api.get("/api/clientes/resumen/").status_code, 403)

    def test_resumen_no_tiene_n_mas_1(self):
        def consultas():
            self.api.force_authenticate(self.oficina)
            with CaptureQueriesContext(connection) as ctx:
                self.assertEqual(self.api.get("/api/clientes/resumen/").status_code, 200)
            return len(ctx)

        for i in range(3):
            self.control(mata(self.a, f"P{i}"))
        pocas = consultas()
        for i in range(40):
            self.control(mata(self.a, f"G{i}"))
            mata(self.b, f"H{i}")
        self.assertEqual(consultas(), pocas)


# ---------------------------------------------------------------------------
# V2.2 · Paso 3: carga masiva desde Excel / CSV
# ---------------------------------------------------------------------------
ENCABEZADO = ("numero_serie", "clase", "ubicacion", "vencimiento_carga", "vencimiento_ph")
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def xlsx_bytes(filas, encabezado=ENCABEZADO, hoja="Equipos"):
    wb = Workbook()
    ws = wb.active
    ws.title = hoja
    ws.append(list(encabezado))
    for f in filas:
        ws.append(list(f))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@override_settings(AXES_ENABLED=False, SECURE_SSL_REDIRECT=False)
class ImportacionMasivaTests(TestCase):
    def setUp(self):
        self.api = APIClient()
        self.a = Cliente.objects.create(nombre="A")
        self.b = Cliente.objects.create(nombre="B")
        self.admin = User.objects.create_superuser("admin", "admin@x.com", "Clave-Segura-123")
        self.oficina = self._user("ofi", Rol.OFICINA, [self.a])
        self.operario = self._user("ope", Rol.OPERARIO, [self.a])
        self.publico = self._user("pub", None, [])
        self.api.force_authenticate(self.oficina)

    def _user(self, username, rol, clientes):
        u = User.objects.create_user(username, f"{username}@x.com", "Clave-Segura-123")
        u.perfil.rol = rol
        u.perfil.save()
        u.perfil.clientes.set(clientes)
        return u

    def subir(self, contenido, nombre="equipos.xlsx", cliente=None, dry_run=False, mime=XLSX_MIME):
        return self.api.post(
            "/api/matafuegos/importar/",
            {
                "archivo": SimpleUploadedFile(nombre, contenido, content_type=mime),
                "cliente": (cliente or self.a).id,
                "dry_run": "true" if dry_run else "false",
            },
            format="multipart",
        )

    def textos_error(self, r):
        return [e["texto"] for e in r.json()["errores"]]

    # -- Plantilla ------------------------------------------------------------------
    def test_plantilla_para_staff(self):
        for u in (self.oficina, self.admin):
            self.api.force_authenticate(u)
            r = self.api.get("/api/matafuegos/plantilla-importacion/")
            self.assertEqual(r.status_code, 200)
            self.assertEqual(r["Content-Type"], XLSX_MIME)
            self.assertIn("plantilla_matafuegos.xlsx", r["Content-Disposition"])
        wb = load_workbook(io.BytesIO(r.content))
        self.assertEqual(wb.sheetnames, ["Equipos", "Instrucciones"])
        self.assertEqual([c.value for c in wb["Equipos"][1]], list(ENCABEZADO))
        self.assertEqual(wb["Equipos"]["A2"].value, "EJEMPLO-001")  # una fila de ejemplo
        self.assertIsNone(wb["Equipos"]["A3"].value)

    def test_plantilla_se_puede_importar_tal_cual_menos_el_ejemplo(self):
        r = self.api.get("/api/matafuegos/plantilla-importacion/")
        r = self.subir(r.content, dry_run=True)
        self.assertEqual(r.status_code, 200, r.content)
        self.assertEqual((r.json()["validas"], r.json()["con_error"]), (1, 0))
        self.assertEqual(r.json()["muestra"][0]["numero_serie"], "EJEMPLO-001")

    def test_plantilla_prohibida_para_operario_publico_y_anonimo(self):
        for u, esperado in ((self.operario, 403), (self.publico, 403), (None, 401)):
            self.api.force_authenticate(u)
            self.assertEqual(self.api.get("/api/matafuegos/plantilla-importacion/").status_code, esperado)

    # -- Importación válida --------------------------------------------------------
    def test_importacion_valida_guarda_historial_y_una_sola_auditoria(self):
        contenido = xlsx_bytes([
            ("N1", "ABC", "Hall", date(2027, 3, 15), "20/04/2030"),  # celda de fecha y texto dd/mm/aaaa
            ("N2", "BC", "Taller", None, None),
            (3, None, None, None, None),  # serie numérica
        ])
        r = self.subir(contenido, nombre="mis equipos.xlsx")
        self.assertEqual(r.status_code, 201, r.content)
        body = r.json()
        self.assertEqual(body["creados"], 3)
        self.assertEqual(len(body["ids"]), 3)

        creados = Matafuego.objects.filter(id__in=body["ids"])
        self.assertEqual(set(creados.values_list("numero_serie", flat=True)), {"N1", "N2", "3"})
        self.assertTrue(all(m.cliente_id == self.a.id and m.activo for m in creados))
        n1 = creados.get(numero_serie="N1")
        self.assertEqual(n1.vencimiento_carga, date(2027, 3, 15))
        self.assertEqual(n1.vencimiento_ph, date(2030, 4, 20))
        self.assertEqual(len({m.token_qr for m in creados}), 3)  # un QR único por equipo

        hist = Matafuego.history.filter(id__in=body["ids"])
        self.assertEqual(hist.count(), 3)
        self.assertTrue(all(h.history_type == "+" and h.history_user_id == self.oficina.id for h in hist))

        eventos = Auditoria.objects.filter(accion="IMPORTACION")
        self.assertEqual(eventos.count(), 1)  # UNA entrada por importación
        ev = eventos.get()
        self.assertEqual((ev.usuario_nombre, ev.rol, ev.cliente_id), ("ofi", "OFICINA", self.a.id))
        self.assertEqual(ev.despues["cantidad"], 3)
        self.assertEqual(ev.despues["archivo"], "mis equipos.xlsx")
        self.assertFalse(Auditoria.objects.filter(accion="CREACION", objeto="Matafuego").exists())

    def test_dry_run_no_guarda_ni_audita(self):
        r = self.subir(xlsx_bytes([("N1", "ABC", "", None, None), ("N2", "", "", "mala", None)]), dry_run=True)
        self.assertEqual(r.status_code, 200)
        b = r.json()
        self.assertEqual((b["dry_run"], b["total_filas"], b["validas"], b["con_error"]), (True, 2, 1, 1))
        self.assertEqual(self.textos_error(r), ["Fila 3: fecha de carga inválida (usá dd/mm/aaaa)"])
        self.assertEqual(Matafuego.objects.count(), 0)
        self.assertFalse(Auditoria.objects.filter(accion="IMPORTACION").exists())

    def test_importa_csv_con_punto_y_coma_y_acentos_cp1252(self):
        csv_txt = "Número de serie;Clase;Ubicación;Venc. carga;Venc. PH\nC1;ABC;Cocina ñandú;15/03/2027;\n"
        r = self.subir(csv_txt.encode("cp1252"), nombre="datos.csv", mime="text/csv")
        self.assertEqual(r.status_code, 201, r.content)
        m = Matafuego.objects.get(numero_serie="C1")
        self.assertEqual((m.ubicacion, m.vencimiento_carga), ("Cocina ñandú", date(2027, 3, 15)))

    def test_mil_filas_con_pocas_consultas(self):
        filas = [(f"S{i:04d}", "ABC", f"Piso {i % 9}", date(2027, 1, 1), None) for i in range(1000)]
        contenido = xlsx_bytes(filas)
        with CaptureQueriesContext(connection) as ctx:
            r = self.subir(contenido)
        self.assertEqual(r.status_code, 201, r.content)
        self.assertEqual(Matafuego.objects.count(), 1000)
        self.assertEqual(Matafuego.history.count(), 1000)
        self.assertLess(len(ctx), 40)  # inserción en bloque, no una consulta por equipo

    # -- Validaciones ------------------------------------------------------------
    def test_duplicados_dentro_del_archivo(self):
        r = self.subir(xlsx_bytes([("D1", "", "", None, None), ("D2", "", "", None, None), ("D1", "", "", None, None)]), dry_run=True)
        self.assertEqual(r.json()["con_error"], 1)
        self.assertIn("Fila 4: n° de serie «D1» repetido en el archivo (ya estaba en la fila 2)", self.textos_error(r))

    def test_duplicados_contra_la_base_incluye_bajas_y_no_guarda_nada(self):
        mata(self.a, "X1")
        baja = mata(self.a, "X2")
        baja.activo = False
        baja.save()
        mata(self.b, "X3")  # misma serie en OTRO cliente: permitida
        r = self.subir(xlsx_bytes([("X1", "", "", None, None), ("X2", "", "", None, None), ("X3", "", "", None, None), ("NUEVO", "", "", None, None)]))
        self.assertEqual(r.status_code, 400)
        errores = self.textos_error(r)
        self.assertEqual(len(errores), 2)
        self.assertIn("Fila 2: n° de serie «X1» ya existe para este cliente", errores)
        self.assertTrue(errores[1].startswith("Fila 3:") and "dado de baja" in errores[1])
        self.assertFalse(Matafuego.objects.filter(numero_serie="NUEVO").exists())  # todo o nada

    def test_fecha_invalida_indica_la_fila_exacta(self):
        filas = [(f"F{i}", "", "", None, None) for i in range(12)] + [("MALA", "", "", "31/02/2027", None)]
        r = self.subir(xlsx_bytes(filas), dry_run=True)
        self.assertEqual(self.textos_error(r), ["Fila 14: fecha de carga inválida (usá dd/mm/aaaa)"])

    def test_fecha_ph_invalida_y_fuera_de_rango(self):
        r = self.subir(xlsx_bytes([("P1", "", "", None, "hoy"), ("P2", "", "", date(1850, 1, 1), None)]), dry_run=True)
        t = self.textos_error(r)
        self.assertIn("Fila 2: fecha de PH inválida (usá dd/mm/aaaa)", t)
        self.assertIn("Fila 3: fecha de carga fuera de rango (usá dd/mm/aaaa)", t)

    def test_serie_obligatoria_y_largos_maximos(self):
        r = self.subir(xlsx_bytes([("", "ABC", "", None, None), ("L1", "C" * 31, "", None, None), ("S" * 61, "", "", None, None), ("U1", "", "u" * 201, None, None)]), dry_run=True)
        t = self.textos_error(r)
        self.assertIn("Fila 2: falta el n° de serie", t)
        self.assertIn("Fila 3: clase demasiado largo (máximo 30 caracteres)", t)
        self.assertIn("Fila 4: n° de serie demasiado largo (máximo 60 caracteres)", t)
        self.assertIn("Fila 5: ubicación demasiado largo (máximo 200 caracteres)", t)

    def test_filas_vacias_se_ignoran_y_conservan_la_numeracion(self):
        r = self.subir(xlsx_bytes([("V1", "", "", None, None), (None,) * 5, (None,) * 5, ("V2", "", "", "x", None)]), dry_run=True)
        self.assertEqual(r.json()["total_filas"], 2)
        self.assertEqual(self.textos_error(r), ["Fila 5: fecha de carga inválida (usá dd/mm/aaaa)"])

    def test_formulas_no_se_evaluan(self):
        r = self.subir(xlsx_bytes([("=1+1", "", "", None, None), ("OK", "", "", "=HOY()", None)]), dry_run=True)
        t = self.textos_error(r)
        self.assertTrue(any(x.startswith("Fila 2:") and "fórmula" in x for x in t), t)
        self.assertTrue(any(x.startswith("Fila 3:") and "fórmula" in x for x in t), t)
        self.assertFalse(Matafuego.objects.exists())

    def test_falta_columna_numero_serie(self):
        r = self.subir(xlsx_bytes([("ABC",)], encabezado=("clase",)))
        self.assertEqual(r.status_code, 400)
        self.assertIn("numero_serie", r.json()["detail"])

    def test_archivo_sin_filas(self):
        self.assertEqual(self.subir(xlsx_bytes([])).status_code, 400)

    # -- Límites y contenido real ----------------------------------------------------
    def test_archivo_de_mas_de_5mb_es_413(self):
        grande = b"numero_serie\n" + b"X" * (5 * 1024 * 1024 + 10)
        r = self.subir(grande, nombre="grande.csv", mime="text/csv")
        self.assertEqual(r.status_code, 413)
        self.assertFalse(Matafuego.objects.exists())

    def test_mas_de_5000_filas_se_rechaza(self):
        csv_txt = "numero_serie\n" + "\n".join(f"S{i}" for i in range(5001))
        r = self.subir(csv_txt.encode(), nombre="muchas.csv", mime="text/csv")
        self.assertEqual(r.status_code, 400)
        self.assertIn("5.000", r.json()["detail"])
        ok = "numero_serie\n" + "\n".join(f"S{i}" for i in range(5000))
        self.assertEqual(self.subir(ok.encode(), nombre="justo.csv", mime="text/csv", dry_run=True).json()["validas"], 5000)

    def test_no_se_confia_en_extension_ni_mime(self):
        # texto plano que dice ser xlsx (con MIME de xlsx)
        self.assertEqual(self.subir(b"numero_serie\nA1\n", nombre="falso.xlsx").status_code, 400)
        # un xlsx real que dice ser csv
        r = self.subir(xlsx_bytes([("A1", "", "", None, None)]), nombre="falso.csv", mime="text/csv")
        self.assertEqual(r.status_code, 400)
        # binario cualquiera como csv
        self.assertEqual(self.subir(b"MZ\x90\x00\x03\x00\x00\x00", nombre="x.csv", mime="text/csv").status_code, 400)
        # extensión no permitida aunque el contenido sea válido
        self.assertEqual(self.subir(xlsx_bytes([("A1", "", "", None, None)]), nombre="datos.txt").status_code, 400)
        self.assertEqual(self.subir(xlsx_bytes([("A1", "", "", None, None)]), nombre="datos.xls").status_code, 400)
        self.assertFalse(Matafuego.objects.exists())

    def _zip(self, archivos):
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
            for nombre, datos in archivos.items():
                z.writestr(nombre, datos)
        return buf.getvalue()

    def test_zip_bomb_y_xml_con_entidades_se_rechazan(self):
        base = {"[Content_Types].xml": "<Types/>", "xl/workbook.xml": "<workbook/>"}
        bomba = self._zip({**base, "xl/relleno.bin": b"\0" * (51 * 1024 * 1024)})  # comprime a pocos KB
        self.assertLess(len(bomba), 5 * 1024 * 1024)
        self.assertEqual(self.subir(bomba).status_code, 400)
        xxe = self._zip({**base, "xl/workbook.xml": '<?xml version="1.0"?><!DOCTYPE x [<!ENTITY a "b">]><workbook/>'})
        self.assertEqual(self.subir(xxe).status_code, 400)
        self.assertEqual(self.subir(self._zip({"cualquiera.txt": "hola"})).status_code, 400)

    # -- Permisos por rol y por cliente -------------------------------------------------
    def test_roles_sin_permiso(self):
        contenido = xlsx_bytes([("P1", "", "", None, None)])
        for u, esperado in ((self.operario, 403), (self.publico, 403), (None, 401)):
            self.api.force_authenticate(u)
            self.assertEqual(self.subir(contenido).status_code, esperado)
        self.assertFalse(Matafuego.objects.exists())

    def test_oficina_solo_importa_a_sus_clientes(self):
        contenido = xlsx_bytes([("P1", "", "", None, None)])
        r = self.subir(contenido, cliente=self.b)  # oficina tiene solo A
        self.assertEqual(r.status_code, 403)
        self.assertFalse(Matafuego.objects.exists())
        self.assertEqual(self.subir(contenido, cliente=self.b, dry_run=True).status_code, 403)  # ni siquiera valida

    def test_cliente_inexistente_responde_igual_que_ajeno(self):
        r = self.api.post(
            "/api/matafuegos/importar/",
            {"archivo": SimpleUploadedFile("a.xlsx", xlsx_bytes([("P1", "", "", None, None)]), content_type=XLSX_MIME), "cliente": 99999},
            format="multipart",
        )
        self.assertEqual(r.status_code, 403)

    def test_admin_importa_a_cualquier_cliente(self):
        self.api.force_authenticate(self.admin)
        r = self.subir(xlsx_bytes([("P1", "", "", None, None)]), cliente=self.b)
        self.assertEqual(r.status_code, 201, r.content)
        self.assertEqual(Matafuego.objects.get(numero_serie="P1").cliente_id, self.b.id)

    def test_requiere_archivo_y_cliente(self):
        self.assertEqual(self.api.post("/api/matafuegos/importar/", {"cliente": self.a.id}, format="multipart").status_code, 400)
        self.assertEqual(self.api.post("/api/matafuegos/importar/", {}, format="multipart").status_code, 400)

    # -- Atomicidad -------------------------------------------------------------------
    def test_un_error_en_la_ultima_fila_no_deja_nada_guardado(self):
        filas = [(f"A{i}", "ABC", "", None, None) for i in range(99)] + [("MALA", "", "", "no-es-fecha", None)]
        r = self.subir(xlsx_bytes(filas))
        self.assertEqual(r.status_code, 400)
        self.assertEqual(r.json()["validas"], 99)
        self.assertEqual(r.json()["errores"][0]["fila"], 101)
        self.assertEqual(Matafuego.objects.count(), 0)
        self.assertEqual(Matafuego.history.count(), 0)
        self.assertFalse(Auditoria.objects.filter(accion="IMPORTACION").exists())

    def test_carrera_con_otra_carga_revierte_todo(self):
        """Si la base rechaza el guardado a mitad de camino (serie duplicada por una carga
        simultánea), la transacción se revierte completa: ni equipos, ni historial, ni auditoría."""

        def falla_a_mitad(objs, model, **kw):
            model.objects.bulk_create(objs[:2])  # ya escribió algo...
            raise IntegrityError("duplicate key")  # ...y luego se cae

        with mock.patch("apps.core.views.bulk_create_with_history", side_effect=falla_a_mitad):
            r = self.subir(xlsx_bytes([(f"R{i}", "", "", None, None) for i in range(5)]))
        self.assertEqual(r.status_code, 409)
        self.assertEqual(Matafuego.objects.count(), 0)
        self.assertFalse(Auditoria.objects.filter(accion="IMPORTACION").exists())


@override_settings(AXES_ENABLED=False, SECURE_SSL_REDIRECT=False)
class QRClienteNoAsignadoTests(TestCase):
    """V2.2 · Paso 5: contrato de status que usa ScanQR.jsx para distinguir los casos.
    404 en /matafuegos/qr/<token>/ = equipo de un cliente NO asignado (aviso de asignación);
    cualquier otro código es otro problema y el front no debe mostrar ese aviso."""

    def setUp(self):
        self.api = APIClient()
        self.a = Cliente.objects.create(nombre="A")
        self.b = Cliente.objects.create(nombre="B")
        self.operario = self._user("ope", Rol.OPERARIO, [self.a])
        self.publico = self._user("pub", None, [])
        self.propio = mata(self.a, "PROPIO")
        self.ajeno = mata(self.b, "AJENO", carga=HOY + timedelta(days=5))
        self.ajeno.ubicacion = "Sala de servidores"
        self.ajeno.save()

    def _user(self, username, rol, clientes):
        u = User.objects.create_user(username, f"{username}@x.com", "Clave-Segura-123")
        u.perfil.rol = rol
        u.perfil.save()
        u.perfil.clientes.set(clientes)
        return u

    def privado(self, m):
        return self.api.get(f"/api/matafuegos/qr/{m.token_qr}/")

    def test_operario_con_cliente_asignado_recibe_la_ficha_completa(self):
        self.api.force_authenticate(self.operario)
        r = self.privado(self.propio)
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.json()["id"], self.propio.id)  # datos para "Nuevo control"
        self.assertIn("problemas", r.json())

    def test_cliente_ajeno_es_404_sin_ningun_dato_del_equipo(self):
        self.api.force_authenticate(self.operario)
        r = self.privado(self.ajeno)
        self.assertEqual(r.status_code, 404)
        cuerpo = r.content.decode()
        for dato in ("AJENO", "Sala de servidores", str(self.ajeno.token_qr)):
            self.assertNotIn(dato, cuerpo)
        self.assertEqual(list(r.json().keys()), ["detail"])  # no se agrega nada a lo que ya era público

    def test_la_ficha_publica_del_mismo_equipo_sigue_igual(self):
        self.api.force_authenticate(self.operario)
        r = self.api.get(f"/api/public/qr/{self.ajeno.token_qr}/")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(
            set(r.json()),
            {"numero_serie", "clase", "ubicacion", "cliente", "vencimiento_carga", "vencimiento_ph",
             "estado_color", "ultimo_control"},
        )

    def test_operario_no_puede_controlar_un_equipo_de_cliente_ajeno(self):
        self.api.force_authenticate(self.operario)
        r = self.api.post("/api/controles/", {"matafuego": self.ajeno.id, "presion": "NORMAL"}, format="json")
        self.assertEqual(r.status_code, 400)
        self.assertFalse(Control.objects.filter(matafuego=self.ajeno).exists())

    def test_los_demas_casos_no_son_404(self):
        # sin sesión -> 401 y con cuenta sin rol -> 403: el front ni siquiera consulta el endpoint privado
        self.api.force_authenticate(None)
        self.assertEqual(self.privado(self.propio).status_code, 401)
        self.assertEqual(self.api.get(f"/api/public/qr/{self.propio.token_qr}/").status_code, 200)
        self.api.force_authenticate(self.publico)
        self.assertEqual(self.privado(self.propio).status_code, 403)


# ---------------------------------------------------------------------------
# V2.3 · Estado REVISADO / NO REVISADO (calculado por mes calendario, hora Argentina)
# ---------------------------------------------------------------------------
ART = ZoneInfo("America/Argentina/Buenos_Aires")


def art(y, mes, d, h=12, mi=0):
    return datetime(y, mes, d, h, mi, tzinfo=ART)


def reloj(momento):
    """Fija 'ahora' (timezone.now y por lo tanto localdate/localtime) sin dependencias extra."""
    return mock.patch("django.utils.timezone.now", return_value=momento)


@override_settings(AXES_ENABLED=False, SECURE_SSL_REDIRECT=False)
class RevisionMensualTests(TestCase):
    def setUp(self):
        self.api = APIClient()
        self.a = Cliente.objects.create(nombre="A")
        self.admin = User.objects.create_superuser("admin", "admin@x.com", "Clave-Segura-123")
        self.operario = User.objects.create_user("ope", "ope@x.com", "Clave-Segura-123")
        self.operario.perfil.rol = Rol.OPERARIO
        self.operario.perfil.save()
        self.operario.perfil.clientes.set([self.a])
        self.oficina = User.objects.create_user("ofi", "ofi@x.com", "Clave-Segura-123")
        self.oficina.perfil.rol = Rol.OFICINA
        self.oficina.perfil.save()
        self.oficina.perfil.clientes.set([self.a])

    def control(self, m, fecha=None, **kw):
        if fecha is not None:
            kw["fecha"] = fecha
        return Control.objects.create(matafuego=m, usuario=self.admin, **kw)

    def revisado(self, m):
        # Igual que las vistas: con prefetch_related("controles")
        return Matafuego.objects.prefetch_related("controles").get(pk=m.pk).revisado_mes

    # -- Regla básica -------------------------------------------------------
    def test_sin_controles_no_esta_revisado(self):
        self.assertFalse(self.revisado(mata(self.a)))

    def test_control_de_este_mes_esta_revisado(self):
        m = mata(self.a)
        self.control(m, art(2026, 10, 3))
        with reloj(art(2026, 10, 15)):
            self.assertTrue(self.revisado(m))

    def test_control_del_mes_anterior_no_cuenta(self):
        m = mata(self.a)
        self.control(m, art(2026, 9, 20))
        with reloj(art(2026, 10, 15)):
            self.assertFalse(self.revisado(m))

    def test_mismo_mes_de_otro_anio_no_cuenta(self):
        m = mata(self.a)
        self.control(m, art(2025, 10, 15))
        with reloj(art(2026, 10, 15)):
            self.assertFalse(self.revisado(m))

    def test_cambio_de_mes_reinicia_solo(self):
        m = mata(self.a)
        self.control(m, art(2026, 10, 31, 10, 0))
        with reloj(art(2026, 10, 31, 18, 0)):
            self.assertTrue(self.revisado(m))
        with reloj(art(2026, 11, 1, 9, 0)):  # avanza el reloj al día 1
            self.assertFalse(self.revisado(m))

    def test_un_control_nuevo_en_el_mes_nuevo_vuelve_a_revisar(self):
        m = mata(self.a)
        self.control(m, art(2026, 10, 31, 10, 0))
        self.control(m, art(2026, 11, 2, 10, 0))
        with reloj(art(2026, 11, 5)):
            self.assertTrue(self.revisado(m))

    # -- Borde horario (Argentina es UTC-3) ----------------------------------------
    def test_borde_horario_fin_de_mes_usa_hora_argentina(self):
        # 31/10 23:30 ART == 01/11 02:30 UTC
        m = mata(self.a)
        f = art(2026, 10, 31, 23, 30)
        self.assertEqual(f.astimezone(dt_timezone.utc).month, 11)  # en UTC ya es noviembre
        self.control(m, f)
        with reloj(art(2026, 10, 31, 23, 45)):  # en UTC también es 01/11
            self.assertTrue(self.revisado(m))
        with reloj(art(2026, 11, 1, 0, 10)):  # recién acá empieza noviembre en Argentina
            self.assertFalse(self.revisado(m))

    def test_borde_horario_control_de_septiembre_noche_no_cuenta_en_octubre(self):
        # 30/09 22:00 ART == 01/10 01:00 UTC: para Argentina sigue siendo septiembre
        m = mata(self.a)
        self.control(m, art(2026, 9, 30, 22, 0))
        with reloj(art(2026, 10, 5)):
            self.assertFalse(self.revisado(m))

    # -- Independiente del semáforo ------------------------------------------------
    def test_control_con_problemas_cuenta_como_revision_y_no_cambia_el_semaforo(self):
        m = mata(self.a)
        self.control(m, art(2026, 10, 3), presion=Control.Presion.BAJA)
        with reloj(art(2026, 10, 15)):
            fresco = Matafuego.objects.prefetch_related("controles").get(pk=m.pk)
            self.assertTrue(fresco.revisado_mes)
            self.assertEqual(fresco.estado_color, "ROJO")  # el semáforo sigue diciendo el estado

    def test_operativo_y_no_revisado_a_la_vez(self):
        m = mata(self.a)
        self.control(m, art(2026, 9, 10))
        with reloj(art(2026, 10, 15)):
            fresco = Matafuego.objects.prefetch_related("controles").get(pk=m.pk)
            self.assertEqual(fresco.estado_color, "VERDE")
            self.assertFalse(fresco.revisado_mes)

    # -- API -----------------------------------------------------------------------
    def listar(self, user):
        self.api.force_authenticate(user)
        r = self.api.get(f"/api/matafuegos/?cliente={self.a.id}")
        self.assertEqual(r.status_code, 200)
        return {x["id"]: x for x in r.json()["results"]}

    def test_crear_control_por_api_deja_el_equipo_revisado(self):
        m = mata(self.a)
        self.assertFalse(self.listar(self.operario)[m.id]["revisado_mes"])
        r = self.api.post("/api/controles/", {"matafuego": m.id, "presion": "NORMAL"}, format="json")
        self.assertEqual(r.status_code, 201, r.content)
        self.assertTrue(self.listar(self.operario)[m.id]["revisado_mes"])

    def test_la_fecha_del_control_no_se_puede_falsear_por_api(self):
        m = mata(self.a)
        self.api.force_authenticate(self.operario)
        vieja = "2020-01-01T10:00:00-03:00"
        r = self.api.post(
            "/api/controles/", {"matafuego": m.id, "fecha": vieja}, format="json"
        )
        self.assertEqual(r.status_code, 201, r.content)
        # el servidor ignora la fecha enviada y pone la actual
        self.assertTrue(Control.objects.get(pk=r.json()["id"]).fecha.year >= 2026)
        self.assertTrue(self.listar(self.operario)[m.id]["revisado_mes"])

    def test_qr_publico_no_expone_revisado_mes(self):
        m = mata(self.a)
        self.control(m)
        r = self.api.get(f"/api/public/qr/{m.token_qr}/")
        self.assertEqual(r.status_code, 200)
        self.assertNotIn("revisado_mes", r.json())

    def test_qr_privado_si_incluye_revisado_mes(self):
        m = mata(self.a)
        self.api.force_authenticate(self.operario)
        r = self.api.get(f"/api/matafuegos/qr/{m.token_qr}/")
        self.assertEqual(r.status_code, 200)
        self.assertIn("revisado_mes", r.json())

    def test_listado_no_hace_consultas_por_equipo(self):
        def consultas():
            self.api.force_authenticate(self.oficina)
            with CaptureQueriesContext(connection) as ctx:
                self.assertEqual(self.api.get(f"/api/matafuegos/?cliente={self.a.id}").status_code, 200)
            return len(ctx)

        for i in range(3):
            self.control(mata(self.a, f"P{i}"))
        pocas = consultas()
        for i in range(40):
            self.control(mata(self.a, f"G{i}"))
            mata(self.a, f"H{i}")  # sin controles
        self.assertEqual(consultas(), pocas)

    # -- Resumen por cliente (sin_revisar) ----------------------------------------------
    def test_resumen_cuenta_sin_revisar_solo_activos(self):
        revisado, viejo, nunca = mata(self.a, "R"), mata(self.a, "V"), mata(self.a, "N")
        baja = mata(self.a, "B")
        baja.activo = False
        baja.save()
        self.control(revisado, art(2026, 10, 3))
        self.control(viejo, art(2026, 9, 3))
        self.api.force_authenticate(self.oficina)
        with reloj(art(2026, 10, 15)):
            r = self.api.get("/api/clientes/resumen/")
        self.assertEqual(r.status_code, 200)
        a = {c["id"]: c for c in r.json()}[self.a.id]
        self.assertEqual(a["total"], 3)  # la baja no cuenta
        self.assertEqual(a["sin_revisar"], 2)  # viejo y nunca; la baja tampoco

    def test_resumen_sigue_con_cantidad_fija_de_consultas(self):
        def consultas():
            self.api.force_authenticate(self.oficina)
            with CaptureQueriesContext(connection) as ctx:
                self.assertEqual(self.api.get("/api/clientes/resumen/").status_code, 200)
            return len(ctx)

        for i in range(3):
            self.control(mata(self.a, f"P{i}"))
        pocas = consultas()
        for i in range(30):
            self.control(mata(self.a, f"G{i}"))
        self.assertEqual(consultas(), pocas)
