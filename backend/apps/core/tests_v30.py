"""Tests de la v3.0: borrado desde /admin/ (is_staff) y planilla de extintores por cliente."""
import io
from datetime import date, timedelta

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from . import planilla
from .models import Auditoria, Cliente, Control, Matafuego, Rol

User = get_user_model()
HOY = timezone.localdate()
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MODEL_BACKEND = "django.contrib.auth.backends.ModelBackend"


def usuario(nombre, rol=None, clientes=(), **kw):
    u = User.objects.create_user(nombre, f"{nombre}@x.com", "Clave-Segura-123", **kw)
    u.perfil.rol = rol
    u.perfil.save()
    u.perfil.clientes.set(clientes)
    return u


def equipo(cliente, serie, ubicacion="", **kw):
    kw.setdefault("vencimiento_carga", HOY + timedelta(days=200))
    kw.setdefault("vencimiento_ph", HOY + timedelta(days=200))
    return Matafuego.objects.create(cliente=cliente, numero_serie=serie, ubicacion=ubicacion, **kw)


# ---------------------------------------------------------------------------
# PARTE A · borrar desde /admin/
# ---------------------------------------------------------------------------
@override_settings(AXES_ENABLED=False, SECURE_SSL_REDIRECT=False)
class BorradoAdminTests(TestCase):
    def setUp(self):
        self.cliente = Cliente.objects.create(nombre="ACME")
        self.mata = equipo(self.cliente, "A1", "Hall")
        self.cargador = usuario("cargador", Rol.OPERARIO, [self.cliente])
        self.c1 = Control.objects.create(
            matafuego=self.mata, usuario=self.cargador, observaciones="ok", precinto=True,
            presion=Control.Presion.BAJA, ubicacion="Pasillo",
        )
        self.c2 = Control.objects.create(matafuego=self.mata, usuario=self.cargador)
        # is_staff SIN permisos de modelo y SIN rol en la app: es el caso que tiene que funcionar
        self.staff = usuario("staff", None, is_staff=True)
        self.superuser = User.objects.create_superuser("root", "root@x.com", "Clave-Segura-123")
        # ADMIN de la app pero sin Staff status: no puede
        self.admin_app = usuario("adminapp", Rol.ADMIN, [self.cliente])
        self.web = self.client

    def entrar(self, user):
        self.web.force_login(user, backend=MODEL_BACKEND)

    def borrar(self, control):
        url = reverse("admin:core_control_delete", args=[control.pk])
        return self.web.post(url, {"post": "yes"})

    def test_staff_sin_permisos_de_modelo_borra_un_control(self):
        self.entrar(self.staff)
        r = self.web.get(reverse("admin:core_control_delete", args=[self.c1.pk]))
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "Queda registrado en Auditoría")
        r = self.borrar(self.c1)
        self.assertEqual(r.status_code, 302)
        self.assertFalse(Control.objects.filter(pk=self.c1.pk).exists())
        self.assertTrue(Control.objects.filter(pk=self.c2.pk).exists())

    def test_superusuario_tambien_borra(self):
        self.entrar(self.superuser)
        self.assertEqual(self.borrar(self.c1).status_code, 302)
        self.assertFalse(Control.objects.filter(pk=self.c1.pk).exists())

    def test_admin_de_la_app_sin_staff_no_puede(self):
        self.entrar(self.admin_app)
        r = self.borrar(self.c1)
        self.assertIn(r.status_code, (302, 403))  # redirige al login del admin
        self.assertTrue(Control.objects.filter(pk=self.c1.pk).exists())
        self.assertFalse(Auditoria.objects.filter(accion="ELIMINACION").exists())

    def test_staff_inactivo_no_puede(self):
        self.entrar(self.staff)
        self.staff.is_active = False
        self.staff.save()
        self.borrar(self.c1)
        self.assertTrue(Control.objects.filter(pk=self.c1.pk).exists())

    def test_el_borrado_deja_auditoria_con_los_datos_de_antes(self):
        self.entrar(self.staff)
        self.borrar(self.c1)
        ev = Auditoria.objects.get(accion="ELIMINACION", objeto="Control")
        self.assertEqual(ev.usuario, self.staff)
        self.assertEqual(ev.cliente_id, self.cliente.id)
        self.assertEqual(ev.objeto_id, str(self.c1.pk))
        self.assertEqual(ev.antes["matafuego"], "A1")
        self.assertEqual(ev.antes["cargado_por"], "cargador")
        self.assertEqual(ev.antes["presion"], "BAJA")
        self.assertEqual(ev.antes["observaciones"], "ok")
        self.assertIs(ev.antes["precinto"], True)
        self.assertEqual(ev.antes["ubicacion"], "Pasillo")

    def test_accion_masiva_audita_cada_control(self):
        self.entrar(self.staff)
        r = self.web.post(
            reverse("admin:core_control_changelist"),
            {"action": "delete_selected", "_selected_action": [self.c1.pk, self.c2.pk], "post": "yes"},
        )
        self.assertEqual(r.status_code, 302)
        self.assertEqual(Control.objects.count(), 0)
        self.assertEqual(Auditoria.objects.filter(accion="ELIMINACION", objeto="Control").count(), 2)

    def test_la_accion_masiva_no_existe_para_quien_no_es_staff(self):
        self.entrar(self.admin_app)
        self.web.post(
            reverse("admin:core_control_changelist"),
            {"action": "delete_selected", "_selected_action": [self.c1.pk], "post": "yes"},
        )
        self.assertEqual(Control.objects.count(), 2)

    def test_el_control_sigue_siendo_inmutable_en_el_admin(self):
        self.entrar(self.superuser)
        self.assertEqual(self.web.get(reverse("admin:core_control_add")).status_code, 403)
        self.assertEqual(self.web.get(reverse("admin:core_control_changelist")).status_code, 200)

    def test_delete_directo_sigue_bloqueado_y_la_api_da_405(self):
        with self.assertRaises(PermissionError):
            self.c1.delete()
        with self.assertRaises(PermissionError):
            Control.objects.filter(pk=self.c1.pk).first().delete()
        api = APIClient()
        for u in (self.superuser, self.admin_app, self.staff, self.cargador,
                  usuario("ofi", Rol.OFICINA, [self.cliente])):
            api.force_authenticate(u)
            self.assertEqual(api.delete(f"/api/controles/{self.c1.pk}/").status_code, 405)
            self.assertEqual(api.put(f"/api/controles/{self.c1.pk}/", {}).status_code, 405)
            self.assertEqual(api.patch(f"/api/controles/{self.c1.pk}/", {}).status_code, 405)

    def test_si_falla_la_auditoria_no_se_borra(self):
        from unittest import mock

        with mock.patch("apps.core.audit.auditar", side_effect=RuntimeError("boom")):
            with self.assertRaises(RuntimeError):
                self.c1.eliminar_definitivo(self.staff)
        self.assertTrue(Control.objects.filter(pk=self.c1.pk).exists())

    def test_equipo_y_cliente_se_borran_recien_sin_controles(self):
        self.entrar(self.staff)
        # con controles: el admin lo impide (on_delete=PROTECT)
        self.web.post(reverse("admin:core_matafuego_delete", args=[self.mata.pk]), {"post": "yes"})
        self.assertTrue(Matafuego.objects.filter(pk=self.mata.pk).exists())
        self.borrar(self.c1)
        self.borrar(self.c2)
        # el cliente todavía tiene el matafuego
        self.web.post(reverse("admin:core_cliente_delete", args=[self.cliente.pk]), {"post": "yes"})
        self.assertTrue(Cliente.objects.filter(pk=self.cliente.pk).exists())
        r = self.web.post(reverse("admin:core_matafuego_delete", args=[self.mata.pk]), {"post": "yes"})
        self.assertEqual(r.status_code, 302)
        self.assertFalse(Matafuego.objects.filter(pk=self.mata.pk).exists())
        r = self.web.post(reverse("admin:core_cliente_delete", args=[self.cliente.pk]), {"post": "yes"})
        self.assertEqual(r.status_code, 302)
        self.assertFalse(Cliente.objects.filter(pk=self.cliente.pk).exists())
        self.assertTrue(Auditoria.objects.filter(accion="ELIMINACION", objeto="Matafuego").exists())
        self.assertTrue(Auditoria.objects.filter(accion="ELIMINACION", objeto="Cliente").exists())

    def test_columna_cant_controles_en_el_admin_de_equipos(self):
        self.entrar(self.staff)
        r = self.web.get(reverse("admin:core_matafuego_changelist"))
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "Controles")

    def test_equipo_sin_controles_queda_gris_y_nada_falla(self):
        self.entrar(self.staff)
        self.borrar(self.c1)
        self.borrar(self.c2)
        m = Matafuego.objects.prefetch_related("controles").get(pk=self.mata.pk)
        self.assertEqual(m.estado_color, "GRIS")
        self.assertFalse(m.revisado_mes)
        api = APIClient()
        api.force_authenticate(self.cargador)
        self.assertEqual(api.get("/api/dashboard/vencimientos/").status_code, 200)
        self.assertEqual(api.get("/api/notificaciones/").status_code, 200)
        self.assertEqual(api.get("/api/clientes/resumen/").status_code, 200)
        r = APIClient().get(f"/api/public/qr/{m.token_qr}/")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.json()["estado_color"], "GRIS")


# ---------------------------------------------------------------------------
# PARTE B · campos nuevos
# ---------------------------------------------------------------------------
@override_settings(AXES_ENABLED=False, SECURE_SSL_REDIRECT=False)
class CamposNuevosTests(TestCase):
    def setUp(self):
        self.api = APIClient()
        self.cliente = Cliente.objects.create(nombre="ACME")
        self.ofi = usuario("ofi", Rol.OFICINA, [self.cliente])
        self.api.force_authenticate(self.ofi)

    def test_agente_se_guarda_en_mayusculas(self):
        r = self.api.post("/api/matafuegos/", {
            "cliente": self.cliente.id, "numero_serie": "N1", "marca": "KIDDE",
            "anio_fabricacion": 2020, "agente": "  polvo ", "capacidad": "5 KG",
        }, format="json")
        self.assertEqual(r.status_code, 201, r.content)
        self.assertEqual(r.json()["agente"], "POLVO")
        self.assertEqual(r.json()["capacidad"], "5 KG")

    def test_anio_fuera_de_rango_se_rechaza(self):
        for anio in (1949, HOY.year + 2):
            r = self.api.post("/api/matafuegos/", {
                "cliente": self.cliente.id, "numero_serie": f"N{anio}", "anio_fabricacion": anio,
            }, format="json")
            self.assertEqual(r.status_code, 400)

    def test_control_con_y_sin_precinto(self):
        m = equipo(self.cliente, "P1")
        r = self.api.post("/api/controles/", {"matafuego": m.id, "precinto": False}, format="json")
        self.assertEqual(r.status_code, 201, r.content)
        self.assertIs(r.json()["precinto"], False)
        r = self.api.post("/api/controles/", {"matafuego": m.id}, format="json")
        self.assertIsNone(r.json()["precinto"])  # sin dato: como los controles anteriores a la V3.0

    def test_qr_publico_no_expone_los_campos_nuevos(self):
        m = equipo(self.cliente, "Q1", marca="KIDDE", agente="POLVO", capacidad="5 KG", anio_fabricacion=2020)
        data = APIClient().get(f"/api/public/qr/{m.token_qr}/").json()
        for campo in ("marca", "agente", "capacidad", "anio_fabricacion", "planta", "precinto"):
            self.assertNotIn(campo, data)

    def test_planta_del_cliente(self):
        r = self.api.get(f"/api/clientes/{self.cliente.id}/")
        self.assertEqual(r.status_code, 200)
        self.assertIn("planta", r.json())


# ---------------------------------------------------------------------------
# PARTE B · planilla por cliente
# ---------------------------------------------------------------------------
@override_settings(AXES_ENABLED=False, SECURE_SSL_REDIRECT=False)
class PlanillaTests(TestCase):
    def setUp(self):
        self.api = APIClient()
        self.a = Cliente.objects.create(nombre="Tasa Casa Central", planta="PLANTA HAITI 1371")
        self.b = Cliente.objects.create(nombre="Otro SA")
        self.admin = User.objects.create_superuser("admin", "admin@x.com", "Clave-Segura-123")
        self.ofi_a = usuario("ofi_a", Rol.OFICINA, [self.a])
        self.ofi_b = usuario("ofi_b", Rol.OFICINA, [self.b])
        self.operario = usuario("ope", Rol.OPERARIO, [self.a])
        self.publico = usuario("pub", None)
        self.api.force_authenticate(self.ofi_a)

    def url(self, cliente=None, qs=""):
        return f"/api/clientes/{(cliente or self.a).id}/exportar-planilla/{qs}"

    def descargar(self, cliente=None, qs="?mes=10&anio=2026"):
        r = self.api.get(self.url(cliente, qs))
        self.assertEqual(r.status_code, 200, getattr(r, "content", b"")[:300])
        return r, load_workbook(io.BytesIO(r.content)).worksheets[0]

    def datos_completos(self):
        equipo(self.a, "S-02", "Pasillo", marca="KIDDE", anio_fabricacion=2020, clase="ABC",
               agente="POLVO", capacidad="5 KG",
               vencimiento_ph=date(2030, 3, 15), vencimiento_carga=date(2027, 1, 20))
        equipo(self.a, "S-01", "Hall", marca="YPF", anio_fabricacion=2019, clase="BC",
               agente="CO2", capacidad="10 KG")
        equipo(self.b, "AJENO", "Otro lado")
        equipo(self.a, "BAJA", "Zzz", activo=False)

    # -- Contenido ------------------------------------------------------------------
    def test_conserva_formato_de_la_planilla_base(self):
        base = load_workbook(planilla.PLANTILLA).worksheets[0]
        self.datos_completos()
        _, ws = self.descargar()
        self.assertEqual(sorted(map(str, ws.merged_cells.ranges)), sorted(map(str, base.merged_cells.ranges)))
        self.assertEqual(len(ws._images), len(base._images))
        self.assertGreaterEqual(len(ws._images), 1)  # logo
        for letra in "ABCDEFGHIJKLMN":
            self.assertEqual(ws.column_dimensions[letra].width, base.column_dimensions[letra].width)
        self.assertEqual([c.value for c in ws[3]][:14], [c.value for c in base[3]][:14])
        self.assertEqual(ws.page_setup.orientation, "landscape")

    def test_encabezado_y_mapeo_por_columna(self):
        self.datos_completos()
        _, ws = self.descargar()
        self.assertEqual(ws["C1"].value, "PLANTA HAITI 1371")
        self.assertEqual(ws["D2"].value, planilla.EMPRESA_RELEVADORA)
        self.assertEqual(ws["F1"].value, 2026)
        self.assertEqual(ws["F2"].value, "OCTUBRE")
        # orden por ubicación (Hall < Pasillo) y solo activos del cliente
        self.assertEqual([ws.cell(row=r, column=3).value for r in (4, 5)], ["S-01", "S-02"])
        self.assertEqual([ws.cell(row=r, column=1).value for r in (4, 5)], [1, 2])
        self.assertIsNone(ws["C6"].value)  # las filas de ejemplo no aparecen
        fila = [c.value for c in ws[5]][:14]
        self.assertEqual(fila[1:8], ["Pasillo", "S-02", "KIDDE", 2020, "ABC", "POLVO", "5 KG"])
        self.assertEqual(fila[11].date(), date(2030, 3, 15))  # Vto. PH: fecha real
        self.assertEqual(fila[12].date(), date(2027, 1, 20))  # Vto. Carga
        self.assertEqual(ws["L5"].number_format, "[$-C0A]mmm\\-yy")  # español, como la referencia
        self.assertEqual(ws["L5"].alignment.horizontal, "center")
        self.assertNotIn("AJENO", [c.value for row in ws.iter_rows() for c in row])
        self.assertNotIn("BAJA", [c.value for row in ws.iter_rows() for c in row])

    def test_sin_planta_usa_el_nombre_del_cliente(self):
        self.a.planta = ""
        self.a.save()
        _, ws = self.descargar()
        self.assertEqual(ws["C1"].value, "Tasa Casa Central")

    def test_ok_no_y_vacio_segun_el_ultimo_control(self):
        con = equipo(self.a, "C1", "A")
        equipo(self.a, "SIN", "B")
        equipo(self.a, "VIEJO", "C")
        Control.objects.create(matafuego=con, usuario=self.ofi_a, precinto=True, accesible=False,
                               senalizacion=True, presion=Control.Presion.BAJA,
                               observaciones="Golpeado\ncon marca")
        viejo = Matafuego.objects.get(numero_serie="VIEJO")
        Control.objects.create(matafuego=viejo, usuario=self.ofi_a)  # precinto None (sin dato)
        _, ws = self.descargar()
        filas = {ws.cell(row=r, column=3).value: [c.value for c in ws[r]] for r in (4, 5, 6)}
        self.assertEqual(filas["C1"][8:11], ["OK", "NO OK", "OK"])
        self.assertEqual(filas["SIN"][8:11], [None, None, None])  # sin control: vacío
        self.assertEqual(filas["VIEJO"][8:11], [None, "OK", "OK"])  # precinto sin dato: vacío
        self.assertIn("Golpeado con marca", filas["C1"][13])
        self.assertIn("Presión baja", filas["C1"][13])
        self.assertIn(" · ", filas["C1"][13])
        self.assertNotIn("Inaccesible", filas["C1"][13])  # ya sale como NO en Acceso

    def test_vencimientos_vencidos_van_a_observaciones(self):
        equipo(self.a, "V1", "A", vencimiento_carga=HOY - timedelta(days=5), vencimiento_ph=HOY - timedelta(days=5))
        _, ws = self.descargar()
        self.assertIn("Carga vencida", ws["N4"].value)
        self.assertIn("PH vencida", ws["N4"].value)

    def test_mas_de_12_equipos_extiende_estilos_filtro_y_validaciones(self):
        for i in range(30):
            equipo(self.a, f"E{i:02d}", f"U{i:02d}")
        _, ws = self.descargar()
        self.assertEqual(ws.cell(row=33, column=3).value, "E29")
        self.assertIsNone(ws.cell(row=34, column=3).value)
        self.assertEqual(ws.auto_filter.ref, "A3:N33")
        self.assertEqual(ws.cell(row=30, column=3).border.left.style, ws.cell(row=4, column=3).border.left.style)
        self.assertEqual(ws.row_dimensions[30].height, 24.75)
        rangos = [str(r) for dv in ws.data_validations.dataValidation for r in dv.sqref.ranges]
        self.assertIn("I4:K33", rangos)
        self.assertEqual(ws.print_title_rows, "$1:$3")

    def test_validaciones_son_listas_inline(self):
        _, ws = self.descargar()
        dvs = {str(dv.sqref): dv.formula1 for dv in ws.data_validations.dataValidation}
        self.assertEqual(dvs["F4:F23"], '"ABC,BC,A,K"')
        self.assertEqual(dvs["G4:G23"], '"POLVO,HCFC,CO2,AGUA"')
        self.assertEqual(dvs["H4:H23"], '"5 KG,10 KG"')
        self.assertEqual(dvs["I4:K23"], '"OK,NO OK"')
        self.assertTrue(dvs["E4:E23"].startswith(f'"{HOY.year + 1},{HOY.year},'))
        self.assertNotIn("$AB$", "".join(str(f) for f in dvs.values()))

    def test_texto_con_igual_se_guarda_como_texto_no_formula(self):
        equipo(self.a, "=1+1", '=HYPERLINK("http://malo")', marca="+cmd", agente="@SUM(A1)")
        r, ws = self.descargar()
        self.assertEqual(ws["B4"].data_type, "s")
        self.assertEqual(ws["B4"].value, '=HYPERLINK("http://malo")')
        self.assertEqual(ws["C4"].data_type, "s")
        self.assertEqual(ws["D4"].value, "+cmd")
        self.assertNotIn(b"<f>", __import__("zipfile").ZipFile(io.BytesIO(r.content)).read("xl/worksheets/sheet1.xml"))

    def test_cliente_sin_equipos_genera_planilla_vacia(self):
        _, ws = self.descargar(self.a)
        self.assertIsNone(ws["C4"].value)

    # -- Respuesta, nombre de archivo y auditoría -------------------------------------
    def test_encabezados_nombre_saneado_y_auditoria(self):
        self.datos_completos()
        r, _ = self.descargar()
        self.assertEqual(r["Content-Type"], XLSX_MIME)
        self.assertEqual(r["Content-Disposition"], 'attachment; filename="Extintores_TASA_CASA_CENTRAL_OCTUBRE_2026.xlsx"')
        ev = Auditoria.objects.get(accion="EXPORTACION")
        self.assertEqual(ev.cliente_id, self.a.id)
        self.assertEqual(ev.despues, {"mes": 10, "anio": 2026, "equipos": 2})

    def test_nombre_de_archivo_sin_caracteres_raros(self):
        c = Cliente(nombre='Ñandú "S.A."\r\nX/../y;z')
        nombre = planilla.nombre_archivo(c, 3, 2027)
        self.assertRegex(nombre, r"^Extintores_[A-Z0-9_]+_MARZO_2027\.xlsx$")
        self.assertNotRegex(nombre, r'[\r\n"/;\\]')

    def test_mes_y_anio_invalidos_dan_400(self):
        for qs in ("?mes=13", "?mes=0", "?mes=abc", "?anio=1999", "?anio=2101", "?anio=x"):
            self.assertEqual(self.api.get(self.url(qs=qs)).status_code, 400, qs)
        self.assertFalse(Auditoria.objects.filter(accion="EXPORTACION").exists())

    def test_por_defecto_usa_mes_y_anio_actuales(self):
        _, ws = self.descargar(qs="")
        self.assertEqual(ws["F1"].value, HOY.year)
        self.assertEqual(ws["F2"].value, planilla.MESES[HOY.month - 1])

    def test_limite_de_filas(self):
        from unittest import mock

        with mock.patch.object(planilla, "MAX_FILAS", 1):
            equipo(self.a, "L1")
            equipo(self.a, "L2")
            self.assertEqual(self.api.get(self.url()).status_code, 400)

    # -- Permisos y scoping -----------------------------------------------------------
    def test_permisos_y_scoping(self):
        casos = (
            (self.admin, self.a, 200), (self.admin, self.b, 200),
            (self.ofi_a, self.a, 200), (self.ofi_a, self.b, 404),  # cliente ajeno: no se revela
            (self.ofi_b, self.a, 404),
            (self.operario, self.a, 403), (self.publico, self.a, 403), (None, self.a, 401),
        )
        for user, cliente, esperado in casos:
            self.api.force_authenticate(user)
            r = self.api.get(self.url(cliente, "?mes=10&anio=2026"))
            self.assertEqual(r.status_code, esperado, f"{user} -> {cliente}")
            r = self.api.get(self.url(cliente) + "resumen/")
            self.assertEqual(r.status_code, esperado, f"resumen {user} -> {cliente}")

    def test_resumen(self):
        equipo(self.a, "R1", "A", marca="K", anio_fabricacion=2020, agente="CO2", capacidad="5 KG")
        equipo(self.a, "R2", "B")  # incompleto
        ok = Matafuego.objects.get(numero_serie="R1")
        Control.objects.create(matafuego=ok, usuario=self.ofi_a)
        equipo(self.a, "R3", "C", activo=False)
        equipo(self.b, "R4")
        r = self.api.get(self.url() + "resumen/")
        self.assertEqual(r.json(), {"equipos": 2, "incompletos": 1, "sin_control": 1})
        self.assertFalse(Auditoria.objects.filter(accion="EXPORTACION").exists())


# ---------------------------------------------------------------------------
# PARTE B · importación de la propia planilla
# ---------------------------------------------------------------------------
@override_settings(AXES_ENABLED=False, SECURE_SSL_REDIRECT=False)
class ImportarPlanillaTests(TestCase):
    def setUp(self):
        self.api = APIClient()
        self.a = Cliente.objects.create(nombre="A")
        self.b = Cliente.objects.create(nombre="B")
        self.admin = User.objects.create_superuser("admin", "admin@x.com", "Clave-Segura-123")
        self.api.force_authenticate(self.admin)

    def subir(self, contenido, cliente=None, dry_run=True):
        return self.api.post("/api/matafuegos/importar/", {
            "archivo": SimpleUploadedFile("planilla.xlsx", contenido, content_type=XLSX_MIME),
            "cliente": (cliente or self.b).id,
            "dry_run": "true" if dry_run else "false",
        }, format="multipart")

    def planilla_con(self, **celdas):
        """Planilla base (encabezado en la fila 3) con las celdas indicadas."""
        wb = load_workbook(planilla.PLANTILLA)
        ws = wb.worksheets[0]
        for fila in range(4, 14):
            for col in range(1, 15):
                ws.cell(row=fila, column=col).value = None
        for ref, valor in celdas.items():
            ws[ref] = valor
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def test_lee_la_planilla_con_encabezado_en_la_fila_3(self):
        contenido = self.planilla_con(
            A4=1, B4="Hall", C4="S-001", D4="KIDDE", E4=2020, F4="ABC", G4="polvo", H4="5 KG",
            L4=date(2030, 3, 17), M4=date(2027, 1, 20), A5=2, C5="S-002", O6="RETIRO:", A6=3,
        )
        r = self.subir(contenido)
        self.assertEqual(r.status_code, 200, r.content)
        j = r.json()
        self.assertEqual((j["validas"], j["con_error"], j["total_filas"]), (2, 0, 2))
        m = j["muestra"][0]
        self.assertEqual((m["marca"], m["agente"], m["capacidad"], m["anio_fabricacion"]), ("KIDDE", "POLVO", "5 KG", 2020))
        self.assertEqual(m["vencimiento_ph"], "2030-03-01")  # formato mes-año: día 1

    def test_fecha_de_1931_es_error_de_fila(self):
        contenido = self.planilla_con(C4="S-001", L4=date(1931, 7, 1), C5="S-002", M5=date(2026, 12, 1))
        j = self.subir(contenido).json()
        self.assertEqual(j["con_error"], 1)
        self.assertEqual(j["errores"][0]["texto"], "Fila 4: fecha de PH fuera de rango (usá dd/mm/aaaa)")

    def test_fecha_muy_en_el_futuro_es_error(self):
        contenido = self.planilla_con(C4="S-001", M4=date(HOY.year + 16, 1, 1))
        self.assertEqual(self.subir(contenido).json()["con_error"], 1)

    def test_planilla_exportada_se_vuelve_a_importar(self):
        equipo(self.a, "X1", "Hall", marca="KIDDE", anio_fabricacion=2020, agente="POLVO", capacidad="5 KG",
               vencimiento_ph=date(2030, 3, 15), vencimiento_carga=date(2027, 1, 20))
        equipo(self.a, "X2", "Pasillo")
        r = self.api.get(f"/api/clientes/{self.a.id}/exportar-planilla/?mes=10&anio=2026")
        j = self.subir(r.content, cliente=self.b).json()
        self.assertEqual((j["validas"], j["con_error"]), (2, 0))
        self.assertEqual(self.subir(r.content, cliente=self.b, dry_run=False).status_code, 201)
        m = Matafuego.objects.get(cliente=self.b, numero_serie="X1")
        self.assertEqual((m.marca, m.agente, m.anio_fabricacion), ("KIDDE", "POLVO", 2020))
        self.assertEqual(m.vencimiento_ph, date(2030, 3, 1))

    def test_alias_del_typo_de_la_referencia(self):
        wb = load_workbook(planilla.PLANTILLA)
        ws = wb.worksheets[0]
        ws["G3"] = "Agente Extintintor"
        ws["C4"], ws["G4"] = "T1", "agua"
        buf = io.BytesIO()
        wb.save(buf)
        self.assertEqual(self.subir(buf.getvalue()).json()["muestra"][0]["agente"], "AGUA")

    def test_formato_viejo_de_5_columnas_sigue_funcionando(self):
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.append(["numero_serie", "clase", "ubicacion", "vencimiento_carga", "vencimiento_ph"])
        ws.append(["V1", "ABC", "Hall", date(2027, 3, 15), date(2030, 3, 15)])
        buf = io.BytesIO()
        wb.save(buf)
        j = self.subir(buf.getvalue()).json()
        self.assertEqual((j["validas"], j["con_error"]), (1, 0))
        self.assertEqual(j["muestra"][0]["vencimiento_carga"], "2027-03-15")

    def test_anio_de_fabricacion_invalido(self):
        j = self.subir(self.planilla_con(C4="S-1", E4=1900)).json()
        self.assertIn("año de fabricación", j["errores"][0]["texto"])
