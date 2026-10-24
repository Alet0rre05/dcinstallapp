"""Planilla de extintores por cliente (V3.0).

Se genera a partir de la planilla de referencia guardada en
`plantillas/planilla_extintores.xlsx` (logo, celdas combinadas, anchos, bordes, fuentes,
orientación y «ajustar a una página» se conservan): se abre la base, se limpian las filas de
ejemplo y se completan los datos del cliente. NO se reconstruye el formato celda por celda.

Seguridad: todo texto que viene de la base se escribe como string literal; nunca se evalúa.
"""
import io
import os
import re
import unicodedata
from copy import copy
from pathlib import Path

from django.utils import timezone
from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

from .models import Matafuego

PLANTILLA = Path(__file__).resolve().parent / "plantillas" / "planilla_extintores.xlsx"
EMPRESA_RELEVADORA = os.environ.get("EMPRESA_RELEVADORA", "DCINSTALL SRL")
MAX_FILAS = 5000  # igual que la importación

MESES = (
    "ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO",
    "JULIO", "AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE",
)
FILA_TITULOS = 3
PRIMERA_FILA = 4
ULTIMA_FILA_EJEMPLO = 13  # si alguien vuelve a copiar la referencia original, sus filas 4 a 13 traen ejemplos
ULTIMA_FILA_CON_FORMATO = 22  # filas 4 a 22 con bordes; la 23 va sin bordes (más allá se clona el estilo de la 4)
ULTIMA_FILA_FILTRO = 23  # el filtro cubre A3:N23 como mínimo
ALTO_FILA = 24.75
COLUMNAS_DATOS = 14  # A..N
TEXTO_PLANTA_EN_BLANCO = "PLANTA"
# Listas desplegables (V3.2). Son listas en línea: la referencia apuntaba a celdas auxiliares vacías.
LISTA_CLASES = "ABC,BC,A,K"
LISTA_AGENTES = "POLVO,HCFC,CO2,AGUA"
LISTA_CAPACIDADES = "5 KG,10 KG"
LISTA_OK = "OK,NO OK"
ANIO_MIN_LISTA = 1990
# Del Matafuego.problemas() solo van a Observaciones los que no tienen columna propia
# (Acceso y Señalización ya salen como OK/NO).
PROBLEMAS_EN_OBSERVACIONES = ("Presión", "Chapa/baliza", "Carga vencida", "PH vencida")


class ErrorPlanilla(Exception):
    """Pedido inválido (mes/año, demasiados equipos…). La vista lo devuelve como 400."""


def periodo(mes=None, anio=None):
    """(mes, anio) validados; por defecto el mes/año actuales en hora de Buenos Aires."""
    hoy = timezone.localdate()
    try:
        mes = hoy.month if mes in (None, "") else int(mes)
        anio = hoy.year if anio in (None, "") else int(anio)
    except (TypeError, ValueError):
        raise ErrorPlanilla("El mes y el año tienen que ser números.")
    if not 1 <= mes <= 12:
        raise ErrorPlanilla("El mes tiene que estar entre 1 y 12.")
    if not 2000 <= anio <= 2100:
        raise ErrorPlanilla("El año tiene que estar entre 2000 y 2100.")
    return mes, anio


def nombre_archivo(cliente, mes, anio):
    """Extintores_<CLIENTE>_<MES>_<AAAA>.xlsx, sin caracteres raros ni saltos de línea."""
    t = unicodedata.normalize("NFKD", cliente.nombre or "").encode("ascii", "ignore").decode()
    t = re.sub(r"[^A-Za-z0-9]+", "_", t).strip("_")[:60] or "cliente"
    return f"Extintores_{t.upper()}_{MESES[mes - 1]}_{anio}.xlsx"


def equipos_del_cliente(cliente):
    """Solo activos, ordenados por ubicación y n° de serie."""
    return list(
        Matafuego.objects.filter(cliente=cliente, activo=True)
        .order_by("ubicacion", "numero_serie")
        .prefetch_related("controles")
    )


def resumen(cliente):
    """Datos para la vista previa: equipos activos, con datos incompletos y sin ningún control."""
    equipos = equipos_del_cliente(cliente)
    incompletos = sum(
        1 for m in equipos if not (m.marca and m.anio_fabricacion and m.agente and m.capacidad)
    )
    sin_control = sum(1 for m in equipos if m.ultimo_control is None)
    return {"equipos": len(equipos), "incompletos": incompletos, "sin_control": sin_control}


# ---------------------------------------------------------------------------
def _escribir_texto(celda, valor):
    """Texto de la base como string literal: si empieza con = + - @ nunca es fórmula."""
    texto = "" if valor is None else str(valor)
    celda.value = texto if texto else None
    if texto[:1] in ("=", "+", "-", "@", "\t", "\r"):
        celda.data_type = "s"
        celda.quotePrefix = True


def _ok(valor):
    return "" if valor is None else ("OK" if valor else "NO OK")


def _observaciones(m):
    partes = []
    c = m.ultimo_control
    if c and c.observaciones.strip():
        partes.append(re.sub(r"\s+", " ", c.observaciones).strip())
    partes += [p for p in m.problemas() if p.startswith(PROBLEMAS_EN_OBSERVACIONES)]
    return " · ".join(partes)


def _fila_de(m, n):
    c = m.ultimo_control
    return [
        n, m.ubicacion, m.numero_serie, m.marca, m.anio_fabricacion, m.clase, m.agente, m.capacidad,
        _ok(c.precinto) if c else "", _ok(c.accesible) if c else "", _ok(c.senalizacion) if c else "",
        m.vencimiento_ph, m.vencimiento_carga, _observaciones(m),
    ]


def _extender_area_impresion(ws, ultima):
    area = ws.print_area
    if isinstance(area, (list, tuple)):
        area = area[0] if area else None
    if area:
        m = re.search(r"\$(\d+)$", area)
        if m and int(m.group(1)) < ultima:
            ws.print_area = re.sub(r"\$(\d+)$", f"${ultima}", area.split("!")[-1])
    else:
        ws.print_area = f"A1:O{ultima}"


def _lista(formula, rango, aviso):
    """Lista desplegable en línea. Es un AVISO (no bloquea): si el equipo trae otro valor se puede dejar."""
    dv = DataValidation(
        type="list", formula1=f'"{formula}"', allow_blank=True,
        showErrorMessage=True, errorStyle="warning", errorTitle="Valor fuera de la lista", error=aviso,
    )
    dv.add(rango)
    return dv


def _lista_anios():
    tope = timezone.localdate().year + 1
    return ",".join(str(a) for a in range(tope, ANIO_MIN_LISTA - 1, -1))  # 38 años ≈ 190 caracteres (< 255)


def _armar(planta, mes, anio, equipos):
    """Abre la planilla de referencia y la completa. NO toca fuentes, bordes, anchos ni alineaciones."""
    wb = load_workbook(PLANTILLA)
    ws = wb.worksheets[0]

    # Encabezado (mismo formato de la referencia; solo cambia el texto)
    _escribir_texto(ws["C1"], planta)
    _escribir_texto(ws["D2"], EMPRESA_RELEVADORA)
    ws["F1"] = anio
    ws["F2"] = MESES[mes - 1]

    # Filas de ejemplo fuera
    for fila in range(PRIMERA_FILA, ULTIMA_FILA_EJEMPLO + 1):
        for col in range(1, COLUMNAS_DATOS + 1):
            ws.cell(row=fila, column=col).value = None

    # Datos. Donde la referencia ya no trae formato (fila 23 en adelante) se clona el estilo de la fila 4.
    modelo = [ws.cell(row=PRIMERA_FILA, column=col) for col in range(1, COLUMNAS_DATOS + 2)]  # A..O
    for i, m in enumerate(equipos):
        fila = PRIMERA_FILA + i
        if fila > ULTIMA_FILA_CON_FORMATO:
            for col, origen in enumerate(modelo, start=1):
                ws.cell(row=fila, column=col)._style = copy(origen._style)
            ws.row_dimensions[fila].height = ALTO_FILA
        for col, valor in enumerate(_fila_de(m, i + 1), start=1):
            celda = ws.cell(row=fila, column=col)
            if isinstance(valor, str):
                _escribir_texto(celda, valor)
            else:
                celda.value = valor  # Vto. PH / Vto. Carga: fecha real; el formato «mmm-yy» ya lo trae la celda

    ultima = max(ULTIMA_FILA_FILTRO, PRIMERA_FILA + len(equipos) - 1)

    # Listas desplegables
    ws.data_validations.dataValidation = []
    for dv in (
        _lista(_lista_anios(), f"E{PRIMERA_FILA}:E{ultima}", "El año no está en la lista."),
        _lista(LISTA_CLASES, f"F{PRIMERA_FILA}:F{ultima}", "La clase no está en la lista."),
        _lista(LISTA_AGENTES, f"G{PRIMERA_FILA}:G{ultima}", "El agente no está en la lista."),
        _lista(LISTA_CAPACIDADES, f"H{PRIMERA_FILA}:H{ultima}", "La capacidad no está en la lista."),
        _lista(LISTA_OK, f"I{PRIMERA_FILA}:K{ultima}", "Elegí OK o NO OK."),
    ):
        ws.add_data_validation(dv)

    ws.auto_filter.ref = f"A{FILA_TITULOS}:N{ultima}"
    _extender_area_impresion(ws, ultima)
    ws.print_title_rows = f"1:{FILA_TITULOS}"

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def generar(cliente, mes=None, anio=None, equipos=None):
    """Devuelve (bytes del .xlsx, cantidad de equipos)."""
    mes, anio = periodo(mes, anio)
    equipos = equipos_del_cliente(cliente) if equipos is None else equipos
    if len(equipos) > MAX_FILAS:
        raise ErrorPlanilla(f"El cliente tiene {len(equipos)} equipos activos y el máximo por planilla es {MAX_FILAS}.")
    return _armar(cliente.planta or cliente.nombre, mes, anio, equipos), len(equipos)


def generar_en_blanco(mes=None, anio=None):
    """La misma planilla sin equipos: es la plantilla que se descarga en Importar (V3.2)."""
    mes, anio = periodo(mes, anio)
    return _armar(TEXTO_PLANTA_EN_BLANCO, mes, anio, [])
