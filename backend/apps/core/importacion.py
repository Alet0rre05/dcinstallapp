"""Carga masiva de matafuegos desde Excel (.xlsx) o CSV (V2.2).

Este módulo solo LEE y VALIDA: no toca la base de datos más que para consultar
duplicados. Guardar (bulk_create, historial, auditoría) es responsabilidad de la vista.

Seguridad:
- La extensión no alcanza: se verifica el contenido real (firma ZIP + estructura de
  un .xlsx; texto sin bytes nulos en el .csv).
- Se acota el tamaño descomprimido (zip bombs) y se rechazan XML con DOCTYPE/ENTITY.
- openpyxl se usa en modo read_only y NUNCA se evalúan fórmulas: una celda con
  fórmula se informa como error de fila, no se calcula.
"""
import csv
import io
import os
import re
import unicodedata
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta

from django.utils import timezone
from openpyxl import load_workbook

from .models import Matafuego

MAX_BYTES = 5 * 1024 * 1024
MAX_FILAS = 5000
MAX_FILAS_FISICAS = 100_000  # filas vacías con formato que Excel suele arrastrar
MAX_DESCOMPRIMIDO = 50 * 1024 * 1024
MAX_ERRORES_VISIBLES = 200
LOTE_CONSULTA = 500

COLUMNAS = (
    "numero_serie", "clase", "ubicacion", "marca", "anio_fabricacion", "agente", "capacidad",
    "vencimiento_carga", "vencimiento_ph",
)
FILAS_BUSQUEDA_ENCABEZADO = 10  # la planilla de extintores trae el encabezado en la fila 3
ANIO_MIN_VENCIMIENTO = 2000
ANIOS_MAX_VENCIMIENTO = 15  # una fecha más allá de hoy + 15 años es casi seguro un error de carga

# Encabezados aceptados (ya normalizados: minúsculas, sin tildes, con "_").
# Los de la planilla de extintores («Nro de Serie», «Año Fabric.», «Vto. PH»…) entran como alias;
# «Puesto Nro.», «Prescinto», «Acceso», «Señalización», «Observaciones» y columnas de relleno se ignoran.
ALIAS = {
    "numero_serie": {"numero_serie", "n_serie", "nro_serie", "serie", "numero_de_serie", "nro_de_serie", "n_de_serie"},
    "clase": {"clase"},
    "ubicacion": {"ubicacion"},
    "marca": {"marca"},
    "anio_fabricacion": {"anio_fabricacion", "anio_fabric", "ano_fabricacion", "ano_fabric", "anio", "ano"},
    "agente": {"agente", "agente_extintor", "agente_extintintor"},  # la referencia tiene un typo
    "capacidad": {"capacidad"},
    "vencimiento_carga": {"vencimiento_carga", "venc_carga", "vto_carga"},
    "vencimiento_ph": {"vencimiento_ph", "venc_ph", "vto_ph"},
}
ETIQUETA_FECHA = {"vencimiento_carga": "fecha de carga", "vencimiento_ph": "fecha de PH"}
FORMATOS_FECHA = ("%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%Y-%m-%d")
MESES_ABREV = {
    "ene": 1, "feb": 2, "mar": 3, "abr": 4, "may": 5, "jun": 6,
    "jul": 7, "ago": 8, "sep": 9, "set": 9, "oct": 10, "nov": 11, "dic": 12,
}


class ErrorArchivo(Exception):
    """Problema del archivo en sí (no de una fila). `status` es el HTTP a devolver."""

    def __init__(self, mensaje, status=400):
        super().__init__(mensaje)
        self.status = status


@dataclass
class Resultado:
    archivo: str = ""
    total_filas: int = 0
    filas: list = field(default_factory=list)  # filas válidas: dict listo para Matafuego(...)
    errores: list = field(default_factory=list)  # {"fila", "campo", "mensaje", "texto"}

    @property
    def con_error(self):
        return len({e["fila"] for e in self.errores})

    @property
    def validas(self):
        return len(self.filas)


# ---------------------------------------------------------------------------
# Lectura segura del archivo
# ---------------------------------------------------------------------------
def nombre_seguro(archivo):
    nombre = os.path.basename(archivo.name or "archivo")
    nombre = re.sub(r"[\x00-\x1f\x7f]", "", nombre).strip()
    return nombre[:200] or "archivo"


def _normalizar_encabezado(valor):
    t = unicodedata.normalize("NFD", str(valor or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "_", t.lower()).strip("_")


def _verificar_xlsx(archivo):
    archivo.seek(0)
    if archivo.read(4) != b"PK\x03\x04":
        raise ErrorArchivo("El archivo no es un Excel (.xlsx) válido. Si es un .xls viejo, guardalo como .xlsx.")
    archivo.seek(0)
    if not zipfile.is_zipfile(archivo):
        raise ErrorArchivo("El archivo no es un Excel (.xlsx) válido.")
    archivo.seek(0)
    try:
        with zipfile.ZipFile(archivo) as z:
            infos = z.infolist()
            nombres = {i.filename for i in infos}
            if "xl/workbook.xml" not in nombres or "[Content_Types].xml" not in nombres:
                raise ErrorArchivo("El archivo no tiene la estructura de un Excel (.xlsx).")
            if sum(i.file_size for i in infos) > MAX_DESCOMPRIMIDO:
                raise ErrorArchivo("El archivo es demasiado grande una vez descomprimido.")
            for i in infos:
                if i.filename.endswith((".xml", ".rels")):
                    with z.open(i) as f:
                        inicio = f.read(4096).lower()
                    if b"<!doctype" in inicio or b"<!entity" in inicio:
                        raise ErrorArchivo("El archivo contiene contenido no permitido.")
    except zipfile.BadZipFile:
        raise ErrorArchivo("El archivo está dañado o no es un Excel (.xlsx) válido.")
    archivo.seek(0)


def _es_formato_mes_anio(formato):
    """True para formatos de fecha SIN día (mmm-yy, mm/yyyy…): ahí la persona solo ve mes y año."""
    f = re.sub(r'\[[^\]]*\]|"[^"]*"|\\.', "", str(formato or "")).lower()
    return "m" in f and "y" in f and "d" not in f


def _valor_celda(celda):
    valor = getattr(celda, "value", None)
    if isinstance(valor, (datetime, date)) and _es_formato_mes_anio(getattr(celda, "number_format", "")):
        return date(valor.year, valor.month, 1)  # mes-año -> día 1 de ese mes
    return valor


def _filas_xlsx(archivo):
    _verificar_xlsx(archivo)
    try:
        wb = load_workbook(archivo, read_only=True, data_only=False)
    except Exception:
        raise ErrorArchivo("No se pudo leer el Excel. Verificá que no esté dañado ni protegido con contraseña.")
    try:
        ws = wb["Equipos"] if "Equipos" in wb.sheetnames else wb.worksheets[0]
        for n, fila in enumerate(ws.iter_rows(), start=1):
            if n > MAX_FILAS_FISICAS:
                raise ErrorArchivo(f"El archivo tiene demasiadas filas (máximo {MAX_FILAS:,}).".replace(",", "."))
            yield n, [_valor_celda(c) for c in fila]
    finally:
        wb.close()


def _filas_csv(archivo):
    archivo.seek(0)
    crudo = archivo.read()
    if crudo.startswith(b"PK\x03\x04"):
        raise ErrorArchivo("El archivo parece un Excel. Subilo con extensión .xlsx o guardalo como CSV.")
    if b"\x00" in crudo:
        raise ErrorArchivo("El archivo no es un CSV de texto válido.")
    for codificacion in ("utf-8-sig", "cp1252"):  # Excel en español exporta en cp1252
        try:
            texto = crudo.decode(codificacion)
            break
        except UnicodeDecodeError:
            continue
    else:
        texto = crudo.decode("latin-1")
    primera = texto.split("\n", 1)[0]
    delimitador = max(";,\t", key=primera.count) if any(d in primera for d in ";,\t") else ","
    try:
        for n, fila in enumerate(csv.reader(io.StringIO(texto), delimiter=delimitador), start=1):
            if n > MAX_FILAS_FISICAS:
                raise ErrorArchivo(f"El archivo tiene demasiadas filas (máximo {MAX_FILAS:,}).".replace(",", "."))
            yield n, fila
    except csv.Error:
        raise ErrorArchivo("No se pudo leer el CSV. Verificá el formato.")


def leer_filas(archivo):
    """Devuelve (nro_fila, [valores]) de las filas con contenido; el nro es el que ve la persona en Excel."""
    nombre = nombre_seguro(archivo)
    ext = nombre.lower().rsplit(".", 1)[-1] if "." in nombre else ""
    if ext not in ("xlsx", "csv"):
        raise ErrorArchivo("Solo se aceptan archivos .xlsx o .csv.")
    if archivo.size > MAX_BYTES:
        raise ErrorArchivo("El archivo supera el máximo de 5 MB.", status=413)
    if archivo.size == 0:
        raise ErrorArchivo("El archivo está vacío.")
    origen = _filas_xlsx(archivo) if ext == "xlsx" else _filas_csv(archivo)
    for n, fila in origen:
        if any(c is not None and str(c).strip() != "" for c in fila):
            yield n, fila


# ---------------------------------------------------------------------------
# Conversión de celdas (todo es dato, nada se evalúa)
# ---------------------------------------------------------------------------
def _largo_max(campo):
    return Matafuego._meta.get_field(campo).max_length


def _texto(valor):
    """Celda -> str limpio. Devuelve (texto, error)."""
    if valor is None:
        return "", None
    if isinstance(valor, (datetime, date)):
        return "", "tiene formato de fecha; formateá la columna como Texto"
    if isinstance(valor, bool):
        valor = str(valor)
    elif isinstance(valor, float) and valor.is_integer():
        valor = str(int(valor))  # 123.0 -> "123"
    else:
        valor = str(valor)
    valor = re.sub(r"\s+", " ", valor).strip()
    if valor.startswith("="):
        return "", "contiene una fórmula (no se evalúan); pegá solo el valor"
    return valor, None


def _fecha_mes_anio(texto):
    """«mar-27», «mar 2027», «03/2027» o «3-27» -> primer día de ese mes (o None si no es mes-año)."""
    m = re.fullmatch(r"([a-záéíóú]{3,10})[\s./-]+(\d{2}|\d{4})", texto.lower())
    if m:
        mes = MESES_ABREV.get(unicodedata.normalize("NFD", m.group(1)).encode("ascii", "ignore").decode()[:3])
        anio = int(m.group(2))
    else:
        m = re.fullmatch(r"(\d{1,2})[./-](\d{2}|\d{4})", texto)
        if not m:
            return None
        mes, anio = int(m.group(1)), int(m.group(2))
    if mes is None or not 1 <= mes <= 12:
        return None
    if anio < 100:
        anio += 2000
    return date(anio, mes, 1)


def _limite_vencimiento():
    hoy = timezone.localdate()
    try:
        return hoy.replace(year=hoy.year + ANIOS_MAX_VENCIMIENTO)
    except ValueError:  # 29/02
        return hoy + timedelta(days=365 * ANIOS_MAX_VENCIMIENTO)


def _fecha(valor):
    """Celda -> date | None. Devuelve (fecha, error)."""
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        return None, None
    if isinstance(valor, datetime):
        f = valor.date()
    elif isinstance(valor, date):
        f = valor
    elif isinstance(valor, str):
        s = valor.strip()
        if s.startswith("="):
            return None, "contiene una fórmula (no se evalúan)"
        f = None
        for formato in FORMATOS_FECHA:
            try:
                f = datetime.strptime(s, formato).date()
                break
            except ValueError:
                continue
        if f is None:
            f = _fecha_mes_anio(s)
        if f is None:
            return None, "inválida"
    else:
        return None, "inválida"
    # Plausibilidad: así se detectan errores de carga tipo «julio de 1931»
    if not ANIO_MIN_VENCIMIENTO <= f.year or f > _limite_vencimiento():
        return None, "fuera de rango"
    return f, None


def _anio(valor):
    """Celda -> año de fabricación (int | None). Devuelve (anio, error)."""
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        return None, None
    if isinstance(valor, (datetime, date, bool)):
        return None, "inválido"
    try:
        n = float(str(valor).strip().replace(",", "."))
    except ValueError:
        return None, "inválido"
    if not n.is_integer():
        return None, "inválido"
    maximo = timezone.localdate().year + 1
    if not 1950 <= int(n) <= maximo:
        return None, f"fuera de rango (entre 1950 y {maximo})"
    return int(n), None


# ---------------------------------------------------------------------------
# Validación de TODO el archivo
# ---------------------------------------------------------------------------
def _error(res, fila, campo, mensaje):
    res.errores.append({"fila": fila, "campo": campo, "mensaje": mensaje, "texto": f"Fila {fila}: {mensaje}"})


def validar(archivo, cliente):
    """Lee y valida el archivo completo. No guarda nada."""
    res = Resultado(archivo=nombre_seguro(archivo))
    filas = leer_filas(archivo)

    # El encabezado puede no estar en la fila 1 (la planilla de extintores lo trae en la 3):
    # se busca en las primeras filas la que tenga la columna del n° de serie.
    pos, hubo_filas = {}, False
    for n, celdas in filas:
        hubo_filas = True
        pos = {}
        for i, titulo in enumerate(celdas):
            clave = _normalizar_encabezado(titulo)
            for campo, alias in ALIAS.items():
                if clave in alias and campo not in pos:
                    pos[campo] = i
        if "numero_serie" in pos:
            break
        if n >= FILAS_BUSQUEDA_ENCABEZADO:
            pos = {}
            break
    if not hubo_filas:
        raise ErrorArchivo("El archivo no tiene filas para importar.")
    if "numero_serie" not in pos:
        raise ErrorArchivo("Falta la columna «numero_serie» (o «Nro de Serie») en las primeras filas. Descargá la plantilla y usala como base.")

    series_vistas = {}  # numero_serie -> primera fila donde aparece
    candidatas = []  # (fila, datos) sin errores propios todavía
    for n, celdas in filas:
        # Filas sin nada en las columnas que se importan (ej. «RETIRO:»/«ENTREGA:» de la planilla
        # o solo el «Puesto Nro.» numerado): no son equipos, se ignoran
        if not any(
            i < len(celdas) and celdas[i] is not None and str(celdas[i]).strip() != ""
            for i in pos.values()
        ):
            continue
        res.total_filas += 1
        if res.total_filas > MAX_FILAS:
            raise ErrorArchivo(f"El archivo supera el máximo de {MAX_FILAS:,} filas.".replace(",", "."))

        def celda(campo, celdas=celdas):
            i = pos.get(campo)
            return celdas[i] if i is not None and i < len(celdas) else None

        datos, hubo_error, errores_antes = {}, False, len(res.errores)
        for campo, etiqueta in (
            ("numero_serie", "n° de serie"), ("clase", "clase"), ("ubicacion", "ubicación"),
            ("marca", "marca"), ("agente", "agente"), ("capacidad", "capacidad"),
        ):
            valor, err = _texto(celda(campo))
            if err:
                _error(res, n, campo, f"{etiqueta} {err}")
                hubo_error = True
            elif len(valor) > _largo_max(campo):
                _error(res, n, campo, f"{etiqueta} demasiado largo (máximo {_largo_max(campo)} caracteres)")
                hubo_error = True
            datos[campo] = valor.upper() if campo == "agente" else valor  # igual que Matafuego.save()
        anio, err = _anio(celda("anio_fabricacion"))
        if err:
            _error(res, n, "anio_fabricacion", f"año de fabricación {err}")
            hubo_error = True
        datos["anio_fabricacion"] = anio
        if not datos["numero_serie"] and len(res.errores) == errores_antes:
            _error(res, n, "numero_serie", "falta el n° de serie")
            hubo_error = True
        for campo in ("vencimiento_carga", "vencimiento_ph"):
            valor, err = _fecha(celda(campo))
            if err:
                _error(res, n, campo, f"{ETIQUETA_FECHA[campo]} {err} (usá dd/mm/aaaa)")
                hubo_error = True
            datos[campo] = valor

        serie = datos["numero_serie"]
        if serie:
            if serie in series_vistas:
                _error(res, n, "numero_serie", f"n° de serie «{serie}» repetido en el archivo (ya estaba en la fila {series_vistas[serie]})")
                hubo_error = True
            else:
                series_vistas[serie] = n
        if not hubo_error:
            candidatas.append((n, datos))

    if res.total_filas == 0:
        raise ErrorArchivo("El archivo no tiene filas para importar.")

    # Duplicados contra la base: la serie es única por cliente (incluye equipos dados de baja).
    existentes = {}
    todas = list(series_vistas)
    for i in range(0, len(todas), LOTE_CONSULTA):
        lote = todas[i : i + LOTE_CONSULTA]
        existentes.update(
            Matafuego.objects.filter(cliente=cliente, numero_serie__in=lote).values_list("numero_serie", "activo")
        )
    for n, datos in candidatas:
        activo = existentes.get(datos["numero_serie"])
        if activo is None:
            res.filas.append(datos)
        else:
            extra = "" if activo else " (está dado de baja: restauralo en lugar de importarlo)"
            _error(res, n, "numero_serie", f"n° de serie «{datos['numero_serie']}» ya existe para este cliente{extra}")
    # Los errores de la base se agregaron al final: se ordenan por fila para leerlos de arriba hacia abajo.
    res.errores.sort(key=lambda e: e["fila"])
    return res


# ---------------------------------------------------------------------------
# Plantilla descargable
# ---------------------------------------------------------------------------
def generar_plantilla():
    """V3.2: la plantilla de importación es la planilla de extintores en blanco (mismo formato que
    la de «Planillas»). El importador ya lee ese formato: encabezado en la fila 3, columnas A–N."""
    from . import planilla  # import tardío: planilla importa modelos, no hay ciclo, pero así queda explícito

    return planilla.generar_en_blanco()
