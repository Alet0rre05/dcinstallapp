"""Tests de la v3.2: la planilla (Planillas e Importar) es réplica exacta del formato de referencia."""
import io

from django.test import SimpleTestCase
from openpyxl import load_workbook

from . import importacion, planilla


def hoja(contenido):
    return load_workbook(io.BytesIO(contenido)).worksheets[0]


class FormatoReplicaTests(SimpleTestCase):
    """Se compara celda por celda con lo pedido (nada de rellenos, fuentes y medidas de la referencia)."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.ws = hoja(importacion.generar_plantilla())

    def test_pestana_pagina_y_filtro(self):
        ws = self.ws
        self.assertEqual(ws.title, "EXTINTORES TASA-CASA CENTRAL")
        self.assertEqual(ws.page_setup.orientation, "landscape")
        self.assertEqual(ws.page_setup.paperSize, 9)  # A4
        for lado in ("left", "right", "top", "bottom"):
            self.assertEqual(getattr(ws.page_margins, lado), 0)
        self.assertIsNone(ws.freeze_panes)
        self.assertEqual(ws.auto_filter.ref, "A3:N23")

    def test_combinadas(self):
        self.assertEqual(
            sorted(map(str, self.ws.merged_cells.ranges)),
            sorted(["B1:B2", "C1:C2", "D1:E1", "D2:E2", "F1:G1", "F2:G2", "H1:N2"]),
        )

    def test_anchos_y_altos(self):
        ws = self.ws
        for letra, ancho in dict(A=5, B=47.86, C=22, D=17.14, E=12.43, F=9, G=10, H=12.57, L=11.43, N=46.86, O=20.57).items():
            self.assertEqual(ws.column_dimensions[letra].width, ancho, letra)
        self.assertEqual(ws.row_dimensions[1].height, 22.5)
        self.assertEqual(ws.row_dimensions[2].height, 25.5)
        self.assertEqual(ws.row_dimensions[3].height, 87)
        for fila in range(4, 24):
            self.assertEqual(ws.row_dimensions[fila].height, 24.75, fila)

    def test_logo_anclado_y_con_su_tamano(self):
        (img,) = self.ws._images
        desde = img.anchor._from
        self.assertEqual((desde.col, desde.row, desde.colOff, desde.rowOff), (0, 0, 257175, 76200))
        self.assertEqual((img.anchor.ext.width, img.anchor.ext.height), (2057400, 504825))  # 216 x 53 px

    def test_encabezado(self):
        ws = self.ws
        for ref, fuente, tam in (("C1", "Arial", 12), ("D1", "Arial", 12), ("F1", "Arial", 12), ("H1", "Arial", 12),
                                 ("D2", "Arial", 11), ("F2", "Arial", 11)):
            self.assertEqual((ws[ref].font.name, ws[ref].font.sz, ws[ref].font.b), (fuente, tam, True), ref)
            self.assertEqual(ws[ref].alignment.horizontal, "center", ref)
        self.assertTrue(ws["C1"].alignment.wrap_text)
        self.assertEqual((ws["O1"].value, ws["O2"].value, ws["O3"].value), ("RETIRO: ", "ENTREGA: ", "NUMERO"))

    def test_titulos_y_rotaciones(self):
        ws = self.ws
        for col in "ABCDEFGHIJKLMN":
            c = ws[f"{col}3"]
            self.assertTrue(c.font.b, col)
            self.assertEqual(c.alignment.text_rotation, 90 if col in "AEFGHIJKLM" else 0, col)
            self.assertEqual(c.font.name, "Calibri" if col == "A" else "Arial", col)
            self.assertEqual(c.font.sz, 12, col)
        self.assertEqual(ws["A3"].value, "Puesto \nNro.")
        self.assertEqual(ws["G3"].value, "Agente \nExtintintor")  # se respeta el texto tal cual

    def test_cuerpo_bordes_fuentes_y_sin_rellenos(self):
        ws = self.ws
        for fila in range(4, 23):
            for col in "ABCDEFGHIJKLMNO":
                c = ws[f"{col}{fila}"]
                self.assertEqual(c.border.left.style, "thin", c.coordinate)
                self.assertEqual(c.border.bottom.style, "thin", c.coordinate)
                self.assertIsNone(c.fill.fill_type, c.coordinate)
            self.assertIsNone(ws[f"P{fila}"].border.left.style)  # la P no lleva bordes
        for col in "ABCDEFGHIJKLMNO":
            self.assertIsNone(ws[f"{col}23"].border.left.style)  # la fila 23 va sin bordes
        self.assertEqual(ws["B4"].alignment.horizontal, "left")
        self.assertEqual((ws["C4"].font.name, ws["C4"].font.sz, ws["C4"].alignment.horizontal), ("Calibri", 12, "center"))
        self.assertEqual((ws["N4"].font.name, ws["N4"].font.sz, ws["N4"].alignment.horizontal), ("Arial", 10, "center"))
        self.assertEqual((ws["O4"].font.name, ws["O4"].font.sz), ("Calibri", 11))
        self.assertIn("mmm", ws["L4"].number_format)
        self.assertIn("C0A", ws["L4"].number_format)  # mes en español

    def test_la_planilla_de_planillas_y_la_de_importar_tienen_el_mismo_formato(self):
        otra = hoja(planilla._armar("PLANTA X", 10, 2026, []))
        for ref in ("A3", "G3", "B4", "N4", "L4", "O1"):
            self.assertEqual(otra[ref].font.name, self.ws[ref].font.name, ref)
            self.assertEqual(otra[ref].alignment.text_rotation, self.ws[ref].alignment.text_rotation, ref)
        self.assertEqual(otra["C1"].value, "PLANTA X")
        self.assertEqual(otra["F1"].value, 2026)
        self.assertEqual(otra["F2"].value, "OCTUBRE")
